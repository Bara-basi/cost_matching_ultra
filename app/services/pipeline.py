"""工作台流水线：把「上传件 / 逐单输入 / 飞书同步」变成一轮拆单 + 成本匹配。

所有数据每次都重新跑（不读历史结果），输入是 PDF / Excel / 一维商品行，
输出是 `outputs/runs/<job>/` 下的一轮结果（拆单明细 + 财务核对行）。
"""
from __future__ import annotations

import json
import sys
import uuid
from collections import Counter
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402

from app.services.parser import parse_declaration  # noqa: E402
from app.services.parser.excel_writer import build_rows as build_parse_rows  # noqa: E402
from app.services.parser.excel_writer import write_workbook  # noqa: E402
from app.services.review import build_rows as build_review_rows  # noqa: E402
from app.services.scope import out_of_scope_reason  # noqa: E402

RUNS_DIR = PROJECT_ROOT / "outputs" / "runs"

# 逐单核验/同步输入的最小列（与 PDF 解析结果的列名一致，便于复用同一条拆单链路）
DECLARATION_COLUMNS = (
    "来源文件", "合同号_1", "报关单号", "商品序号", "报关品名", "海关编码",
    "报关重量", "申报数量", "申报单位", "总价", "币种", "产品类型",
)


def new_run_dir(job_id: str | None = None) -> Path:
    path = RUNS_DIR / (job_id or uuid.uuid4().hex[:12])
    path.mkdir(parents=True, exist_ok=True)
    return path


def write_declarations(rows: list[dict], path: Path) -> int:
    """把一维商品行写成「解析结果」格式的 xlsx，供拆单直接消费。

    写之前先把「财务手工补回的拆单结果」合并回一条报关行：
    同一 `报关单号 + 合同号 + 报关品名` 出现多行时，说明财务把拆单结果又补回了表里
    （一行一条拆分记录），这类行不是真的多商品报关，要把**报关金额与报关重量并回去**，
    当成一条报关行再交给拆单（2026-09-29 用户口径；最终以系统核算结果为准）。
    只作用于飞书/手工输入（PDF 解析结果不走这里），所以真的多商品报关不会被并。
    """
    merged_rows, merged_count = merge_declaration_rows(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "结果"
    sheet.append(list(DECLARATION_COLUMNS))
    kept = 0
    for row in merged_rows:
        contract = str(row.get("合同号_1") or row.get("合同号（应收表格）") or "").strip()
        if not contract or out_of_scope_reason(contract):
            continue
        sheet.append([str(row.get(column) or "") for column in DECLARATION_COLUMNS])
        kept += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    if merged_count:
        print(f"[合并报关行] 把 {merged_count} 条手工拆分行并回了报关行（同报关单号+同品名）", flush=True)
    return kept


def _sum_text(*values) -> str:
    """把若干金额/重量文本相加，输出去掉尾零的字符串（空值当 0）。"""
    from decimal import Decimal, InvalidOperation

    total = Decimal(0)
    for value in values:
        try:
            total += Decimal(str(value).replace(",", "").strip() or "0")
        except (InvalidOperation, AttributeError):
            continue
    text = f"{total:.4f}".rstrip("0").rstrip(".")
    return text or "0"


def merge_declaration_rows(rows: list[dict]) -> tuple[list[dict], int]:
    """同 `报关单号 + 合同号 + 报关品名` 的多行 → 合并成一条（金额/重量相加）。"""
    out: list[dict] = []
    index: dict[tuple[str, str, str], dict] = {}
    merged = 0
    for row in rows:
        declaration = str(row.get("报关单号") or "").strip()
        key = (declaration,
               str(row.get("合同号_1") or row.get("合同号（应收表格）") or "").strip(),
               str(row.get("报关品名") or "").strip())
        if not declaration or key not in index:
            copy = dict(row)
            out.append(copy)
            if declaration:
                index[key] = copy
            continue
        target = index[key]
        target["总价"] = _sum_text(target.get("总价"), row.get("总价"))
        target["报关重量"] = _sum_text(target.get("报关重量"), row.get("报关重量"))
        for name in ("申报数量",):
            if str(row.get(name) or "").strip():
                target[name] = _sum_text(target.get(name), row.get(name))
        target["商品序号"] = ""   # 合并后不再是原来那一条拆分记录
        merged += 1
    return out, merged


def parse_pdfs(pdf_paths: list[Path], out_xlsx: Path) -> dict:
    """解析报关单 PDF → 出口退税联 xlsx（与 `scripts/parse_customs.py` 同口径）。"""
    rows: list[dict] = []
    failed: list[dict] = []
    warnings: Counter = Counter()
    for path in pdf_paths:
        try:
            parsed = parse_declaration(path)
        except Exception as exc:  # noqa: BLE001
            failed.append({"文件": path.name, "原因": str(exc)})
            continue
        warnings.update(parsed.warnings)
        if parsed.header.sheet_type != "出口退税联":
            continue
        for row in build_parse_rows(parsed):
            if out_of_scope_reason(str(row.get("合同号_1") or "")):
                continue
            rows.append(row)
    if rows:
        write_workbook(rows, out_xlsx)
    return {"rows": len(rows), "failed": failed, "warnings": dict(warnings)}


def run_split(parse_xlsx: Path, out_dir: Path, only: str = "") -> dict:
    """调用拆单引擎（脚本的服务入口），返回统计。"""
    from scripts.run_split_knapsack import run as split_run

    return split_run(parse_xlsx=parse_xlsx, out_dir=out_dir, only=only) or {}


def _read_sheet(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    header = [str(cell or "") for cell in rows[0]]
    return [dict(zip(header, values)) for values in rows[1:]]


def _num(value) -> "Decimal":
    from decimal import Decimal, InvalidOperation

    try:
        return Decimal(str(value).replace(",", "").strip() or "0")
    except (InvalidOperation, AttributeError):
        return Decimal(0)


def build_supplements(out_dir: Path, parse_xlsx: Path) -> list[dict]:
    """用睿贝出运明细，给「没有任何报关单认领」的产品行补出补充报关行。

    原则（2026-09-29 用户口径）：**不要求必须有报关单 PDF 才能算成本**。
    拆单跑完后仍未认领的产品行，按「合同号 + 采购单」合并成一条补充报关行，
    金额取这些行的出运金额合计；再把该出运单的差额（客户费用等，
    用 ERP 的「出运总金额」比对）按金额占比摊到这些补充行上。
    """
    from decimal import Decimal

    unclaimed_path = out_dir / "拆单_未认领产品行.xlsx"
    if not unclaimed_path.exists():
        return []
    unclaimed = _read_sheet(unclaimed_path)
    groups: dict[tuple[str, str], dict] = {}
    for row in unclaimed:
        contract = str(row.get("合同号_1") or "").strip()
        purchase = str(row.get("采购单号") or "").strip()
        amount = _num(row.get("出运金额"))
        if not contract or not purchase or amount <= 0:
            continue
        key = (contract, purchase)
        entry = groups.setdefault(key, {"amount": Decimal(0), "name": "", "rows": 0})
        entry["amount"] += amount
        entry["rows"] += 1
        name = str(row.get("报关品名") or "").strip()
        if name and not entry["name"]:
            entry["name"] = name
    if not groups:
        return []

    # 该出运单已有报关行用了什么币种，补充行沿用（金额口径必须一致）
    currency: dict[str, str] = {}
    declared_total: dict[str, Decimal] = {}
    for row in _read_sheet(parse_xlsx):
        contract = str(row.get("合同号_1") or "").strip()
        if not contract:
            continue
        if str(row.get("币种") or "").strip():
            currency.setdefault(contract, str(row["币种"]).strip())
        declared_total[contract] = declared_total.get(contract, Decimal(0)) + _num(row.get("总价"))

    # 每个出运单：ERP 出运总金额 − 已有报关行金额 − 补充行金额 = 还差的（客户费用等）
    supplement_total: dict[str, Decimal] = {}
    for (contract, _purchase), entry in groups.items():
        supplement_total[contract] = supplement_total.get(contract, Decimal(0)) + entry["amount"]
    gap: dict[str, Decimal] = {}
    for contract in supplement_total:
        total, _fee = _shipment_totals(contract)
        if total <= 0:
            continue
        gap[contract] = total - declared_total.get(contract, Decimal(0)) - supplement_total[contract]

    supplements: list[dict] = []
    for (contract, purchase), entry in groups.items():
        amount = entry["amount"]
        extra = Decimal(0)
        share = supplement_total.get(contract, Decimal(0))
        if share > 0 and gap.get(contract):
            extra = (gap[contract] * amount / share).quantize(Decimal("0.01"))
        supplements.append(
            {
                "来源文件": "睿贝出运明细补充",
                "合同号_1": contract,
                "报关单号": "",
                "商品序号": "",
                "报关品名": entry["name"],
                "海关编码": "",
                "报关重量": "",
                "申报数量": "",
                "申报单位": "",
                "总价": f"{amount + extra:.2f}",
                "币种": currency.get(contract, "USD"),
                "产品类型": "",
            }
        )
    return supplements


def append_declarations(parse_xlsx: Path, rows: list[dict]) -> int:
    """把补充报关行追加到解析结果表（同一张 xlsx 上继续跑一轮）。"""
    from openpyxl import load_workbook

    workbook = load_workbook(parse_xlsx)
    sheet = workbook.active
    kept = 0
    for row in rows:
        sheet.append([row.get(column, "") for column in DECLARATION_COLUMNS])
        kept += 1
    workbook.save(parse_xlsx)
    return kept


def run_round(parse_xlsx: Path, out_dir: Path, only: str = "") -> dict:
    """跑一轮拆单；若还有产品行没被任何报关单认领，就用睿贝出运明细补行后**再跑一轮**。"""
    stats = run_split(parse_xlsx, out_dir, only)
    supplements = build_supplements(out_dir, parse_xlsx)
    if not supplements:
        return {"stats": stats, "supplemented": 0}
    append_declarations(parse_xlsx, supplements)
    stats = run_split(parse_xlsx, out_dir, only)
    return {"stats": stats, "supplemented": len(supplements)}


def scan_receivable(rows: list[dict]) -> dict:
    """出运单「客户费用」与「报关单是否齐」的本地核验（需求 3）。

    基准 = 睿贝出运单的 **出运总金额**（已叠加全部客户费用，是最终参考金额）：
    本批同一出运单记录的报关金额合计应当等于它；不等即「这批报关单不齐」。
    若该出运单还带客户费用（罚金 / 运费 / 检验费…），说明金额分配依赖缺失的那几张
    报关单才能确定，此时逐单核验要拒绝并提示用户补齐。

    返回 {出运单/合同号: {"missing": bool, "fee": Decimal, "gap": Decimal, "fee_bad": bool}}
    """
    from decimal import Decimal

    groups: dict[str, dict] = {}
    for row in rows:
        key = str(row.get("合同号_1") or row.get("合同号（应收表格）") or "").strip()
        if not key:
            continue
        entry = groups.setdefault(key, {"declared": Decimal(0), "usd": Decimal(0)})
        entry["declared"] += Decimal(str(row.get("报关金额") or 0).replace(",", "") or 0)
        entry["usd"] += Decimal(str(row.get("出运金额合计") or 0).replace(",", "") or 0)
    if not groups:
        return {}

    out: dict[str, dict] = {}
    for key, entry in groups.items():
        total, fee = _shipment_totals(key)
        if total > 0:
            # 以 ERP 的出运总金额为最终参考（含所有额外费用）
            gap = entry["declared"] - total
        else:
            # 取不到出运总金额时退回「产品金额 + 客户费用」口径
            gap = entry["declared"] - entry["usd"] - fee
        # 金额容差 ±5：一张报关单金额基本不可能低于 50，±5 内视为舍入/汇率误差
        tolerance = Decimal("5")
        missing = abs(gap) > tolerance
        out[key] = {
            "missing": missing,
            "fee": fee,
            "gap": gap,
            "fee_bad": missing and abs(fee) > tolerance,
            "declared": entry["declared"],
            "total": total,
        }
    return out


def _shipment_totals(contract: str) -> tuple["Decimal", "Decimal"]:
    """出运单 → (出运总金额, 客户费用合计)。取不到返回 (0, 0)。"""
    from decimal import Decimal

    from app.services.shipment_detail import load_shipment_detail, strip_pi

    for key in (contract, strip_pi(contract), f"PI-{strip_pi(contract)}"):
        if not key:
            continue
        payload = load_shipment_detail(key)
        if not payload:
            continue
        base = payload.get("baseInfo") or {}
        total = Decimal(str(base.get("出运总金额") or 0).replace(",", "") or 0)
        fee = Decimal(0)
        for item in payload.get("expenseList") or []:
            fee += Decimal(str(item.get("金额") or 0).replace(",", "") or 0)
        return total, fee
    return Decimal(0), Decimal(0)


def review_rows(split_xlsx: Path, amount_source: dict[str, str] | None = None) -> list[dict]:
    return build_review_rows(split_xlsx, amount_source)


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
