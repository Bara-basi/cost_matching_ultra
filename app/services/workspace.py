"""工作台任务：一次「解析 → 拆单 → 成本匹配 → 财务核对行」的编排与状态。

任务目录：`.runtime/jobs/<12位id>/`，状态写在 `status.json`，
结果行写在 `rows.json`，下载文件按类型命名。每次都是重跑，不读历史结果。
"""
from __future__ import annotations

import json
import threading
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.services import erp_sync, pipeline
from app.services.erp_cache import CACHE_ROOT

PROJECT_ROOT = Path(__file__).resolve().parents[2]
JOBS_DIR = PROJECT_ROOT / ".runtime" / "jobs"
RUNS_DIR = PROJECT_ROOT / "outputs" / "runs"

FRIENDLY_COLUMNS = [
    "报关单号", "合同号_1", "合同号（应收表格）", "外销订单号", "产品类型", "报关品名",
    "供应商简称", "采购订单号", "报关金额", "采购金额", "报关重量",
    "异常类型", "异常明细",
]
SOURCE_COLUMNS = ["_睿贝采购单", "_睿贝出运单", "_飞书链接"]

_LOCK = threading.Lock()
_WORKER_LOCK = threading.Lock()
_CANCELLED: set[str] = set()


def _shortage(key: str, entry: dict) -> str:
    """一句给财务看的缺口说明：报关金额 vs 出运总金额、差多少、其中多少是客户费用。"""
    from decimal import Decimal

    def money(value) -> str:
        try:
            return f"{Decimal(str(value or 0)):,.2f}"
        except Exception:  # noqa: BLE001
            return str(value or 0)

    gap = abs(Decimal(str(entry.get("gap") or 0)))
    fee = abs(Decimal(str(entry.get("fee") or 0)))
    total = entry.get("total") or 0
    declared = entry.get("declared") or 0
    if total:
        text = (f"{key}：报关金额合计 {money(declared)}、出运总金额 {money(total)}，"
                f"还差 {money(gap)}")
    else:
        text = f"{key}：金额对不上，还差 {money(gap)}"
    if fee:
        text += f"（其中客户费用 {money(fee)}）"
    return text + "；请补齐对应的报关单再核对。"


def _status_path(job_id: str) -> Path:
    return JOBS_DIR / job_id / "status.json"


def _update(job_id: str, **values) -> dict:
    path = _status_path(job_id)
    with _LOCK:
        try:
            status = json.loads(path.read_text(encoding="utf-8"))
        except Exception:  # noqa: BLE001
            status = {"id": job_id}
        status.update(values)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(status, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return status


def get_job(job_id: str) -> dict | None:
    path = _status_path(job_id)
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except ValueError:
        return None


def cancel_job(job_id: str) -> dict | None:
    status = get_job(job_id)
    if not status:
        return None
    _CANCELLED.add(job_id)
    return _update(job_id, state="cancelled", message="任务已取消")


def result_rows(job_id: str) -> list[dict] | None:
    path = JOBS_DIR / job_id / "rows.json"
    if not path.exists():
        return None
    try:
        rows = json.loads(path.read_text(encoding="utf-8"))
        overrides = JOBS_DIR / job_id / "overrides.json"
        if overrides.exists():
            edits = json.loads(overrides.read_text(encoding="utf-8"))
            for index, values in edits.items():
                if 0 <= int(index) < len(rows):
                    rows[int(index)] = {**rows[int(index)], **values}
        flags_path = JOBS_DIR / job_id / "flags.json"
        if flags_path.exists():
            flags = json.loads(flags_path.read_text(encoding="utf-8"))
            for index, values in flags.items():
                if 0 <= int(index) < len(rows):
                    rows[int(index)] = {**rows[int(index)], **{
                        key: values[key] for key in ("异常类型", "异常明细", "_人工上报")
                        if key in values}}
        reviewed_path = JOBS_DIR / job_id / "reviewed.json"
        if reviewed_path.exists():
            reviewed = json.loads(reviewed_path.read_text(encoding="utf-8"))
            for index in reviewed:
                if 0 <= int(index) < len(rows):
                    rows[int(index)]["_已抽查"] = True
        return rows
    except ValueError:
        return None


def download_path(job_id: str, kind: str) -> Path | None:
    names = {"results": "成本匹配结果.xlsx", "exceptions": "异常记录.xlsx"}
    name = names.get(kind)
    if not name:
        return None
    path = JOBS_DIR / job_id / name
    return path if path.exists() else None


def _write_xlsx(rows: list[dict], path: Path, columns: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "结果"
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    for index, name in enumerate(columns, 1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(12, min(36, len(name) * 2 + 4))
    sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _read_sheet(path: Path) -> list[dict]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        iterator = book.active.iter_rows(values_only=True)
        header = [str(value or "") for value in next(iterator, ())]
        return [dict(zip(header, values)) for values in iterator]
    finally:
        book.close()


def _compare_pdf_evidence(input_rows: list[dict], pdfs: list[Path], run_dir: Path) -> dict:
    out_path = run_dir / "核验附件解析.xlsx"
    stats = pipeline.parse_pdfs(pdfs, out_path)
    if not out_path.exists():
        return {"state": "complete", "rows": [], "message": "附件未解析出出口退税联，暂无法比较", "parse": stats}
    parsed = _read_sheet(out_path)

    def groups(rows):
        grouped: dict[str, dict] = {}
        for row in rows:
            key = str(row.get("报关单号") or "").strip()
            if not key:
                continue
            group = grouped.setdefault(key, {"count": 0, "total": Decimal(0), "contracts": set()})
            group["count"] += 1
            group["contracts"].add(str(row.get("合同号_1") or "").strip())
            try:
                amount = row.get("总价")
                if amount in (None, ""):
                    amount = row.get("报关金额")
                group["total"] += Decimal(str(amount or 0).replace(",", ""))
            except ArithmeticError:
                pass
        return grouped

    original, document = groups(input_rows), groups(parsed)
    results = []
    for declaration, local in original.items():
        other = document.get(declaration)
        if other is None:
            status, reason = "暂无法比较", "附件中未找到对应报关单"
        elif local["count"] == other["count"] and local["contracts"] == other["contracts"] and abs(local["total"] - other["total"]) <= Decimal("1"):
            status, reason = "结果一致", "商品行数、合同号和报关金额合计一致"
        else:
            status, reason = "结果不一致", "商品行数、合同号或报关金额合计存在差异"
        results.append({"报关单号": declaration, "状态": status, "说明": reason,
                        "输入金额": f"{local['total']:.2f}",
                        "附件金额": f"{other['total']:.2f}" if other else ""})
    return {"state": "complete", "rows": results, "parse": stats}


def start_job(
    *,
    kind: str,
    files: list[tuple[str, bytes]] | None = None,
    declarations: list[dict] | None = None,
    use_ai: bool = False,
    source_records: list[dict] | None = None,
) -> str:
    job_id = uuid.uuid4().hex[:12]
    _CANCELLED.discard(job_id)
    job_dir = JOBS_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    uploads = job_dir / "uploads"
    uploads.mkdir(exist_ok=True)
    saved: list[Path] = []
    for name, content in files or []:
        safe = Path(str(name).replace("\\", "/")).name or "upload"
        target = uploads / safe
        index = 1
        while target.exists():
            target = uploads / f"{Path(safe).stem}_{index}{Path(safe).suffix}"
            index += 1
        target.write_bytes(content)
        saved.append(target)
    if source_records is not None:
        (job_dir / "source_records.json").write_text(
            json.dumps(source_records, ensure_ascii=False, default=str), encoding="utf-8"
        )
    if declarations is not None:
        (job_dir / "input_rows.json").write_text(
            json.dumps(declarations, ensure_ascii=False, default=str), encoding="utf-8"
        )
    declaration_numbers = list(dict.fromkeys(str(item.get("报关单号") or "").strip()
                                             for item in declarations or [] if item.get("报关单号")))
    first_contract = next((str(item.get("合同号_1") or "").strip() for item in declarations or []
                           if item.get("合同号_1")), "")
    _update(
        job_id, id=job_id, kind=kind, state="queued", progress=2, message="已接收输入",
        createdAt=datetime.now(timezone.utc).isoformat(),
        fileCount=len(saved), rowCount=len(declarations or []), useAi=use_ai,
        subject=declaration_numbers[0] if declaration_numbers else first_contract,
        declarationCount=len(declaration_numbers),
        downloads={"results": f"/api/jobs/{job_id}/download/results",
                   "exceptions": f"/api/jobs/{job_id}/download/exceptions"},
    )
    thread = threading.Thread(
        target=_run, args=(job_id, kind, saved, declarations or [], use_ai), daemon=True
    )
    thread.start()
    return job_id


def _run(job_id: str, kind: str, files: list[Path], declarations: list[dict], use_ai: bool) -> None:
    with _WORKER_LOCK:
        if job_id in _CANCELLED:
            _update(job_id, state="cancelled", message="任务已取消")
            return
        _run_active(job_id, kind, files, declarations, use_ai)


def _run_active(job_id: str, kind: str, files: list[Path], declarations: list[dict], use_ai: bool) -> None:
    run_dir = pipeline.new_run_dir(job_id)

    def stopped() -> bool:
        return job_id in _CANCELLED

    def halt_if_cancelled() -> bool:
        """取消后立即收尾：把状态钉在 cancelled，绝不被后续 complete 覆盖。"""
        if stopped():
            _update(job_id, state="cancelled", message="任务已取消")
            return True
        return False

    try:
        _update(job_id, state="running", progress=8, message="正在准备报关数据")
        if declarations:
            parse_xlsx = run_dir / "报关输入.xlsx"
            kept = pipeline.write_declarations(declarations, parse_xlsx)
            if not kept:
                raise ValueError("没有可核验的报关商品行（合同号为空或不属于本期范围）")
            parse_stats = {"rows": kept, "failed": [], "warnings": {}}
            evidence_pdfs = [path for path in files if path.suffix.lower() == ".pdf"]
            if evidence_pdfs:
                try:
                    comparison = _compare_pdf_evidence(declarations, evidence_pdfs, run_dir)
                except Exception as exc:
                    comparison = {"state": "complete", "rows": [],
                                  "message": f"附件暂无法比较：{exc}"}
                (JOBS_DIR / job_id / "evidence_comparison.json").write_text(
                    json.dumps(comparison, ensure_ascii=False, default=str), encoding="utf-8")
        else:
            templates = [path for path in files if path.suffix.lower() == ".xlsx"]
            pdfs = [path for path in files if path.suffix.lower() == ".pdf"]
            zips = [path for path in files if path.suffix.lower() == ".zip"]
            for archive in zips:
                import zipfile

                with zipfile.ZipFile(archive) as handle:
                    for member in handle.namelist():
                        if member.lower().endswith(".pdf"):
                            target = run_dir / "uploads" / Path(member).name
                            target.parent.mkdir(parents=True, exist_ok=True)
                            target.write_bytes(handle.read(member))
                            pdfs.append(target)
            if templates and not pdfs and not zips:
                from app.services.template_import import read_template
                rows_in, errors = read_template(templates[0])
                if errors:
                    raise ValueError("模板校验失败：" + "；".join(f"第 {e['row']} 行 {e['message']}" for e in errors[:5]))
                parse_xlsx = run_dir / "报关输入.xlsx"
                kept = pipeline.write_declarations(rows_in, parse_xlsx)
                if not kept:
                    raise ValueError("模板中没有本期可处理的商品行")
                parse_stats = {"rows": kept, "failed": [], "warnings": {}}
            elif not pdfs:
                raise ValueError("没有找到可解析的报关单 PDF")
            else:
                _update(job_id, progress=20, message=f"正在解析 {len(pdfs)} 份报关单")
                parse_xlsx = run_dir / "报关单解析结果_出口退税联.xlsx"
                parse_stats = pipeline.parse_pdfs(pdfs, parse_xlsx)
                if not parse_stats["rows"]:
                    raise ValueError("报关单没有解析出可用商品行，请检查文件")

        if halt_if_cancelled():
            return
        input_path = JOBS_DIR / job_id / "input_rows.json"
        if not input_path.exists() and parse_xlsx.exists():
            book = load_workbook(parse_xlsx, read_only=True, data_only=True)
            try:
                sheet = book.active
                iterator = sheet.iter_rows(values_only=True)
                header = [str(value or "") for value in next(iterator, ())]
                inputs = [dict(zip(header, values)) for values in iterator]
                input_path.write_text(json.dumps(inputs, ensure_ascii=False, default=str), encoding="utf-8")
            finally:
                book.close()
        _update(job_id, progress=45, message="正在拆单（按采购订单 / 供应商 / 产品类型）")
        split_dir = run_dir / "shipments_split"
        # 缺报关单不再挡住计算：没有报关单认领的产品行，用睿贝出运明细补出补充报关行后再跑一轮
        round_result = pipeline.run_round(parse_xlsx, split_dir)
        if halt_if_cancelled():
            return
        stats = round_result.get("stats") or {}
        supplemented = int(round_result.get("supplemented") or 0)
        if supplemented:
            _update(job_id, message=f"已用睿贝出运明细补充 {supplemented} 条报关行后重算")
        split_xlsx = split_dir / "拆单明细_全部.xlsx"
        if not split_xlsx.exists():
            raise ValueError("拆单没有产出结果，请检查合同号是否属于本期范围")

        _update(job_id, progress=70, message="正在匹配采购成本（入库单实发）")
        rows = pipeline.review_rows(split_xlsx)
        if halt_if_cancelled():
            return

        # 出运单金额与报关金额对不上时（例如报关单未到齐），只做记录级提示，不再拒绝计算
        receivable = pipeline.scan_receivable(rows)
        blocked = {key: entry for key, entry in receivable.items()
                   if entry["missing"] and entry["fee_bad"]}
        if blocked:
            for row in rows:
                key = str(row.get("合同号_1") or "")
                entry = blocked.get(key)
                if entry:
                    row["异常类型"] = "出运单有额外费用"
                    row["异常明细"] = _shortage(key, entry)

        _update(job_id, progress=92, message="正在整理结果")
        from app.services.review import summary as review_summary

        payload = review_summary(rows)
        if halt_if_cancelled():
            return
        (JOBS_DIR / job_id / "rows.json").write_text(
            json.dumps(rows, ensure_ascii=False, default=str), encoding="utf-8"
        )
        columns = FRIENDLY_COLUMNS + SOURCE_COLUMNS
        _write_xlsx(rows, JOBS_DIR / job_id / "成本匹配结果.xlsx", columns)
        _write_xlsx(
            [row for row in rows if row.get("异常类型") != "正常"],
            JOBS_DIR / job_id / "异常记录.xlsx",
            columns,
        )
        if halt_if_cancelled():
            return
        _update(
            job_id, state="complete", progress=100, message="处理完成", summary=payload,
            parse=parse_stats, split=stats,
        )
    except Exception as exc:  # noqa: BLE001
        if stopped():
            _update(job_id, state="cancelled", message="任务已取消")
            return
        text = str(exc)
        if len(text) > 300 or "Traceback" in text or ".py" in text:
            text = "处理未完成，请稍后重试；如反复出现请联系管理员"
        _update(job_id, state="failed", message=text, progress=100)


def list_jobs(limit: int = 30) -> list[dict]:
    if not JOBS_DIR.exists():
        return []
    jobs = [item for path in JOBS_DIR.iterdir() if path.is_dir()
            if (item := get_job(path.name))]
    jobs.sort(key=lambda item: item.get("createdAt", ""), reverse=True)
    selected = jobs[:max(1, min(limit, 100))]
    from app.services import feishu_workflow
    for item in selected:
        if item.get("kind") == "sync" and item.get("state") == "complete":
            proposal = feishu_workflow.plan(item["id"])
            push_path = JOBS_DIR / item["id"] / "push_results.json"
            pushed = json.loads(push_path.read_text(encoding="utf-8")) if push_path.exists() else []
            written = {entry.get("sourceId") for entry in pushed if entry.get("status") == "written"}
            item["writebackReady"] = sum(group["status"] == "ready" and group["sourceId"] not in written
                                          for group in (proposal or {}).get("groups", []))
    return selected


def recover_interrupted() -> None:
    """重启后重新排队仍有原始输入的任务；缺输入的旧任务给明确状态。"""
    if not JOBS_DIR.exists():
        return
    for directory in JOBS_DIR.iterdir():
        if not directory.is_dir():
            continue
        status = get_job(directory.name)
        if not status or status.get("state") not in ("queued", "running"):
            continue
        files = list((directory / "uploads").iterdir()) if (directory / "uploads").exists() else []
        input_file = directory / "input_rows.json"
        declarations = (json.loads(input_file.read_text(encoding="utf-8"))
                        if input_file.exists() and status.get("kind") in ("lookup", "sync") else [])
        if not files and not declarations:
            _update(directory.name, state="failed", message="服务重启后缺少原始输入，请重新提交")
            continue
        _update(directory.name, state="queued", progress=2, message="服务重启后重新排队")
        threading.Thread(target=_run, args=(directory.name, status.get("kind", "files"), files,
                                             declarations, bool(status.get("useAi"))), daemon=True).start()


def append_audit(job_id: str, action: str, operator: str, detail: dict) -> None:
    path = JOBS_DIR / job_id / "audit.jsonl"
    entry = {"at": datetime.now(timezone.utc).isoformat(), "action": action,
             "operator": operator, "detail": detail}
    with _LOCK:
        with path.open("a", encoding="utf-8") as output:
            output.write(json.dumps(entry, ensure_ascii=False, default=str) + "\n")


def audit(job_id: str) -> list[dict]:
    path = JOBS_DIR / job_id / "audit.jsonl"
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()] if path.exists() else []


def override_row(job_id: str, index: int, amount: str, operator: str, reason: str = "") -> dict:
    if not operator.strip():
        raise ValueError("请填写操作员称呼")
    rows = result_rows(job_id)
    if rows is None or index < 0 or index >= len(rows):
        raise ValueError("记录不存在")
    value = Decimal(str(amount))
    if value < 0 or value.as_tuple().exponent < -2:
        raise ValueError("采购金额必须为非负数，最多两位小数")
    flags_path = JOBS_DIR / job_id / "flags.json"
    flags = json.loads(flags_path.read_text(encoding="utf-8")) if flags_path.exists() else {}
    if flags.get(str(index), {}).get("state") == "resubmitted":
        raise ValueError("这条记录已重新提交，请在新任务中调整金额")
    path = JOBS_DIR / job_id / "overrides.json"
    edits = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    old = rows[index].get("采购金额", "")
    edits[str(index)] = {"采购金额": f"{value:.2f}", "异常类型": "人工调整待写回",
                         "异常明细": "人工调整采购金额；需在写回预览中确认。", "_人工调整": True}
    path.write_text(json.dumps(edits, ensure_ascii=False, indent=2), encoding="utf-8")
    if str(index) in flags:
        prior_flag = flags.pop(str(index))
        flags_path.write_text(json.dumps(flags, ensure_ascii=False, indent=2), encoding="utf-8")
        append_audit(job_id, "resolve_flag_by_override", operator,
                     {"row": index, "reportedReason": prior_flag.get("异常明细", "")})
    append_audit(job_id, "override", operator, {"row": index, "old": old, "new": str(value), "reason": reason})
    current = result_rows(job_id) or []
    columns = FRIENDLY_COLUMNS + SOURCE_COLUMNS
    _write_xlsx(current, JOBS_DIR / job_id / "成本匹配结果.xlsx", columns)
    _write_xlsx([item for item in current if item.get("异常类型") != "正常"],
                JOBS_DIR / job_id / "异常记录.xlsx", columns)
    from app.services.review import summary as review_summary
    _update(job_id, summary=review_summary(current))
    return edits[str(index)]


def flag_exception(job_id: str, index: int, operator: str, reason: str) -> dict:
    # 与飞书写回共用互斥锁，防止预览通过后、真正落表前插入人工异常。
    from app.services import feishu_workflow
    with feishu_workflow._PUSH_LOCK:
        return _flag_exception_locked(job_id, index, operator, reason)


def _flag_exception_locked(job_id: str, index: int, operator: str, reason: str) -> dict:
    from app.services import feishu_workflow
    status = get_job(job_id)
    if not status or status.get("kind") != "sync" or status.get("state") != "complete":
        raise ValueError("只能上报已完成的飞书区间核算记录")
    if not operator.strip() or not reason.strip():
        raise ValueError("请填写操作员称呼和异常原因")
    rows = result_rows(job_id)
    if rows is None or not 0 <= index < len(rows):
        raise ValueError("结果记录不存在")
    if rows[index].get("异常类型") != "正常":
        raise ValueError("这条记录已是异常，请直接复核")
    proposal = feishu_workflow.plan(job_id)
    group = next((item for item in (proposal or {}).get("groups", [])
                  if any(child["index"] == index for child in item["children"])), None)
    pushed_path = JOBS_DIR / job_id / "push_results.json"
    if group and pushed_path.exists():
        pushed = json.loads(pushed_path.read_text(encoding="utf-8"))
        if any(item.get("sourceId") == group["sourceId"] and item.get("status") == "written"
               for item in pushed):
            raise ValueError("这组记录已经写回，请先在飞书核对，不可再标记为待写回异常")
    path = JOBS_DIR / job_id / "flags.json"
    flags = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    flags[str(index)] = {"异常类型": "人工上报异常", "异常明细": reason.strip(),
                         "_人工上报": True, "state": "active",
                         "reportedAt": datetime.now(timezone.utc).isoformat(),
                         "operator": operator.strip()}
    path.write_text(json.dumps(flags, ensure_ascii=False, indent=2), encoding="utf-8")
    append_audit(job_id, "flag_exception", operator, {"row": index, "reason": reason.strip()})
    current = result_rows(job_id) or []
    _write_xlsx(current, JOBS_DIR / job_id / "成本匹配结果.xlsx", FRIENDLY_COLUMNS + SOURCE_COLUMNS)
    _write_xlsx([item for item in current if item.get("异常类型") != "正常"],
                JOBS_DIR / job_id / "异常记录.xlsx", FRIENDLY_COLUMNS + SOURCE_COLUMNS)
    from app.services.review import summary as review_summary
    _update(job_id, summary=review_summary(current))
    return flags[str(index)]


def all_exceptions() -> list[dict]:
    if not JOBS_DIR.exists():
        return []
    found = []
    statuses = [item for directory in JOBS_DIR.iterdir() if directory.is_dir()
                if (item := get_job(directory.name)) and item.get("state") == "complete"]
    for status in statuses:
        path = JOBS_DIR / status["id"] / "flags.json"
        if not path.exists():
            continue
        flags = json.loads(path.read_text(encoding="utf-8"))
        base_path = JOBS_DIR / status["id"] / "rows.json"
        base_rows = json.loads(base_path.read_text(encoding="utf-8")) if base_path.exists() else []
        for index_text, flag in flags.items():
            index = int(index_text)
            if not 0 <= index < len(base_rows) or not flag.get("_人工上报"):
                continue
            row = base_rows[index]
            found.append({"jobId": status["id"], "createdAt": status.get("createdAt"),
                          "index": index, "declaration": row.get("报关单号"),
                          "contract": row.get("合同号_1"), "product": row.get("报关品名"),
                          "reason": flag.get("异常明细"), "operator": flag.get("operator"),
                          "reportedAt": flag.get("reportedAt"),
                          "state": flag.get("state", "active"), "newJobId": flag.get("newJobId")})
    return sorted(found, key=lambda item: item.get("reportedAt") or item.get("createdAt") or "", reverse=True)


def resubmit_exception(job_id: str, index: int, operator: str) -> dict:
    if not operator.strip():
        raise ValueError("请填写操作员称呼")
    status = get_job(job_id)
    flags_path = JOBS_DIR / job_id / "flags.json"
    if not status or status.get("kind") != "sync" or not flags_path.exists():
        raise ValueError("待重新核验的记录不存在")
    flags = json.loads(flags_path.read_text(encoding="utf-8"))
    flag = flags.get(str(index))
    if not flag or flag.get("state", "active") != "active":
        raise ValueError("这条记录已重新提交或已人工调整")
    input_path = JOBS_DIR / job_id / "input_rows.json"
    source_path = JOBS_DIR / job_id / "source_records.json"
    if not input_path.exists() or not source_path.exists():
        raise ValueError("原任务缺少输入数据，无法重新核验")
    inputs = json.loads(input_path.read_text(encoding="utf-8"))
    sources = json.loads(source_path.read_text(encoding="utf-8"))
    if not inputs or not sources:
        raise ValueError("原任务输入为空，无法重新核验")
    new_id = start_job(kind="sync", declarations=inputs, source_records=sources)
    _update(new_id, selectionFilters=status.get("selectionFilters", {}),
            resubmittedFrom={"jobId": job_id, "index": index})
    flag.update(state="resubmitted", newJobId=new_id,
                resubmittedAt=datetime.now(timezone.utc).isoformat(), resubmittedBy=operator.strip())
    flags_path.write_text(json.dumps(flags, ensure_ascii=False, indent=2), encoding="utf-8")
    append_audit(job_id, "resubmit_exception", operator,
                 {"row": index, "newJobId": new_id, "reason": flag.get("异常明细", "")})
    return {"id": new_id}


def mark_reviewed(job_id: str, index: int, operator: str) -> dict:
    rows = result_rows(job_id)
    if rows is None or not 0 <= index < len(rows):
        raise ValueError("结果记录不存在")
    if not operator.strip():
        raise ValueError("请填写操作员称呼")
    path = JOBS_DIR / job_id / "reviewed.json"
    reviewed = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    reviewed[str(index)] = {"operator": operator, "at": datetime.now(timezone.utc).isoformat()}
    path.write_text(json.dumps(reviewed, ensure_ascii=False, indent=2), encoding="utf-8")
    append_audit(job_id, "sample_review", operator, {"row": index})
    _update(job_id, reviewedRows=len(reviewed))
    return {"reviewed": True, "reviewedRows": len(reviewed)}
