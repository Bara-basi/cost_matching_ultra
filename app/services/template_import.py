"""财务商品行固定模板的生成与校验。"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from decimal import Decimal, InvalidOperation
import json
import re

from openpyxl import Workbook, load_workbook

SCHEMA_PATH = Path(__file__).with_name("feishu_copy_schema.json")
INPUT_COLUMNS = ("报关单号", "合同号_1", "报关品名", "报关金额", "币种", "报关重量")
REQUIRED = INPUT_COLUMNS


def table_columns() -> tuple[str, ...]:
    """下载模板沿用目标飞书副本表的真实列名与顺序。"""
    if SCHEMA_PATH.exists():
        fields = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
        columns = tuple(str(item["field_name"]) for item in fields)
        if all(name in columns for name in INPUT_COLUMNS):
            return columns
    return INPUT_COLUMNS


def template_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "2026年报关数据"
    columns = table_columns()
    sheet.append(columns)
    sheet.freeze_panes = "A2"
    for index, name in enumerate(columns, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = max(16, len(name) * 2 + 4)
        if name in ("报关单号", "合同号_1", "海关编码", "商品序号"):
            for line in range(2, 202):
                sheet.cell(line, index).number_format = "@"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def read_template(path: Path, *, limit: int | None = None) -> tuple[list[dict], list[dict]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        iterator = sheet.iter_rows()
        header = [str(cell.value or "").strip() for cell in next(iterator, ())]
        missing = [name for name in REQUIRED if name not in header]
        if missing:
            return [], [{"row": 1, "message": "缺少必填列：" + "、".join(missing)}]
        rows: list[dict] = []
        errors: list[dict] = []
        for line, cells in enumerate(iterator, 2):
            values = [cell.value for cell in cells]
            if not any(value is not None and str(value).strip() for value in values):
                continue
            row = {name: str(values[i] if i < len(values) and values[i] is not None else "").strip()
                   for i, name in enumerate(header) if name in INPUT_COLUMNS or name == "采购订单号"}
            empty = [name for name in REQUIRED if not row.get(name)]
            if empty:
                errors.append({"row": line, "message": "必填值为空：" + "、".join(empty)})
            elif cells[header.index("报关单号")].data_type == "n":
                errors.append({"row": line, "message": "报关单号必须使用文本格式，以免 Excel 改写长单号"})
            elif not re.fullmatch(r"\d{18}", row["报关单号"]):
                errors.append({"row": line, "message": "报关单号应为 18 位数字文本"})
            else:
                try:
                    if Decimal(row["报关金额"].replace(",", "")) < 0 or Decimal(row["报关重量"].replace(",", "")) < 0:
                        raise InvalidOperation
                except InvalidOperation:
                    errors.append({"row": line, "message": "报关金额和报关重量必须为非负数字"})
                else:
                    row["总价"] = row["报关金额"]
                    row["来源文件"] = path.name
                    rows.append(row)
            if limit and len(rows) > limit:
                errors.append({"row": line, "message": f"单次最多 {limit} 条"})
                break
        if not rows and not errors:
            errors.append({"row": 1, "message": "模板没有商品行"})
        return rows, errors
    finally:
        book.close()
