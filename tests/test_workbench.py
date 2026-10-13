"""工作台的金额守恒、模板和写回授权回归测试。"""
from __future__ import annotations

import json
import tempfile
import unittest
import urllib.error
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app
from app.services import contract_shipments, feishu_workflow, lookup, pipeline, review, shipment_detail, template_import, workspace


class WorkbenchTests(unittest.TestCase):
    def test_missing_grn_batch_never_allocates_other_batch_cost(self):
        from decimal import Decimal as D
        from app.services.cost_match import RecordCost, allocate_amounts
        records = [RecordCost(amount_rmb=D("37960"), batch="A"),
                   RecordCost(amount_rmb=D("506000"), batch="B")]
        costs = allocate_amounts(records, D("38110"), {"A": D("38110")})
        self.assertEqual(costs[0][0], D("38110"))
        self.assertIsNone(costs[1][0])
        self.assertIn("批次成本缺失", costs[1][1])
        both = allocate_amounts(records, D("544110"), {"A": D("38110"), "B": D("506000")})
        self.assertEqual([amount for amount, _ in both], [D("38110"), D("506000")])
        only_b = allocate_amounts(records[1:], D("544110"), {"A": D("38110"), "B": D("506000")})
        self.assertEqual(only_b[0][0], D("506000"))

    def test_combined_grn_reads_blocks_instead_of_title_or_filename_batch(self):
        from app.services.grn_extract import parse_template, blocks_to_amounts
        from app.services.sheet_reader import read_attachment
        from decimal import Decimal as D
        # 合成布局反例：文件名 A、抬头 B，实际同时覆盖 A/B；不修改生产凭证。
        source = Path(".cache/erp/attachments/grn/26MT-03T271-HD/26MT-03T271-A_鸿迪_入库单_回传_.xlsx")
        with tempfile.TemporaryDirectory() as directory:
            book = load_workbook(source, data_only=True)
            sheet = book.active
            sheet.cell(2, 1, "26MT-03T271-B入库单")
            sheet.cell(4, 12, "实发数据-B")
            sheet.cell(5, 12, "数量")
            sheet.cell(5, 13, "实际金额")
            sheet.cell(6, 12, "个")
            sheet.cell(6, 13, "元")
            sheet.cell(8, 12, 2)
            sheet.cell(8, 13, 506000)
            sheet.cell(10, 12, 2)
            sheet.cell(10, 13, 506000)
            path = Path(directory) / "26MT-03T271-A.xlsx"
            book.save(path)
            book.close()
            parsed = parse_template(read_attachment(path))
            self.assertTrue(parsed.ok, parsed.reason)
            self.assertEqual(blocks_to_amounts(parsed.payload["settled_blocks"]), {"A": D("38110"), "B": D("506000")})
            self.assertEqual(D(parsed.payload["settled_amount"]), D("544110"))

    def test_existing_parsed_attachment_does_not_hide_new_batch_file(self):
        from app.services import grn_extract, grn_select
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "grn" / "po"
            parsed = root / "parsed" / "po"
            source.mkdir(parents=True)
            parsed.mkdir(parents=True)
            (source / "A.xlsx").write_bytes(b"old")
            (source / "B.xlsx").write_bytes(b"new")
            (parsed / "A.xlsx.json").write_text(json.dumps({"_digest": __import__("hashlib").sha256(b"old").hexdigest()[:16]}), encoding="utf-8")
            with patch.object(grn_select, "GRN_DIR", root / "grn"), \
                 patch.object(grn_extract, "PARSED_ROOT", root / "parsed"), \
                 patch.object(grn_extract, "extract_attachment") as extract:
                grn_select.ensure_parsed("po")
                extract.assert_called_once_with(source / "B.xlsx", purchase_code="po", use_model=True)
                extract.reset_mock()
                (source / "A.xlsx").write_bytes(b"updated-same-name")
                grn_select.ensure_parsed("po")
                extract.assert_any_call(source / "A.xlsx", purchase_code="po", use_model=True)

    def test_inbound_without_dan_keeps_erp_attachment_version_order(self):
        from app.services import grn_select
        names = ["26MT-03T271-A 入库单.xlsx", "26MT-03T271-B 入库（回传）.xlsx"]
        with tempfile.TemporaryDirectory() as directory, patch.object(grn_select, "LIST_DIR", Path(directory)):
            (Path(directory) / "po.json").write_text(json.dumps({"attachmentList": [
                {"attachmentName": name} for name in names]}), encoding="utf-8")
            self.assertEqual(grn_select.load_order("po"), {names[0]: 0, names[1]: 1})

    def test_unlisted_local_grn_cannot_outrank_current_erp_attachment(self):
        from app.services import grn_select
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory) / "po"
            folder.mkdir()
            for name in ["old.xlsx", "current.xlsx"]:
                (folder / (name + ".json")).write_text(json.dumps({"_file": name}), encoding="utf-8")
            with patch.object(grn_select, "PARSED_ROOT", Path(directory)), \
                 patch.object(grn_select, "ensure_parsed"), \
                 patch.object(grn_select, "load_order", return_value={"current.xlsx": 0}), \
                 patch.object(grn_select, "disk_to_original", return_value={}), \
                 patch.object(grn_select, "erp_supplier", return_value=""):
                items = grn_select.load_items("po")
                self.assertEqual([item.file for item in items], ["old.xlsx", "current.xlsx"])

    def test_unmatched_cost_batch_stays_blocked_after_manual_amount_override(self):
        source = {"record_id": "b", "报关单号": "310120260514919201",
                  "合同号_1": "26MT-03T271B", "报关品名": "不锈钢管件", "报关金额": "87991.48"}
        row = {**source, "采购金额": "38110", "异常类型": "人工调整待写回",
               "_成本未核算": True, "供应商简称": "鸿迪", "采购订单号": "26MT-03T271-HD"}
        with tempfile.TemporaryDirectory() as directory, patch.object(workspace, "JOBS_DIR", Path(directory)):
            job = Path(directory) / "blocked"
            job.mkdir()
            (job / "source_records.json").write_text(json.dumps([source]), encoding="utf-8")
            (job / "rows.json").write_text(json.dumps([row]), encoding="utf-8")
            group = feishu_workflow.plan("blocked")["groups"][0]
            self.assertEqual(group["status"], "blocked")
            self.assertTrue(any("尚未完成核算" in reason for reason in group["reasons"]))

    def test_grn_attachment_names_accept_inbound_without_dan_and_empty_download_keeps_original(self):
        from app.services.erp_attachment_mcp import grn_attachments, download_file
        from unittest.mock import Mock
        names = ["26MT-03T271-B_鸿迪_入库（回传） (1).xlsx", "26MT-03T271-A 入库单.xlsx", "合同.pdf"]
        value = {"attachmentList": [{"attachmentName": name, "downloadUrl": "https://example.test/file"} for name in names]}
        self.assertEqual([item["name"] for item in grn_attachments(value)], names[:2])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cached.xlsx"
            path.write_bytes(b"valid-original")
            page = Mock()
            page.request.get.return_value.status = 200
            page.request.get.return_value.body.return_value = b""
            self.assertEqual(download_file(page, "https://example.test/file", path), 0)
            self.assertEqual(path.read_bytes(), b"valid-original")

    def test_sync_child_failure_not_success_and_download_skip_is_success(self):
        from scripts import erp_sync as sync_script
        from subprocess import CompletedProcess
        with patch.object(sync_script.subprocess, "run", return_value=CompletedProcess([], 0, '{"errors":1}\n', '')), \
             patch.object(sync_script, "write_state"):
            with self.assertRaises(RuntimeError):
                sync_script.run_script(["erp_fetch_grn_mcp.py"])
        with patch.object(sync_script.subprocess, "run", return_value=CompletedProcess([], 0, '{"errors":0,"skipped":3}\n', '')):
            self.assertEqual(sync_script.run_script(["erp_fetch_grn_mcp.py"])["summary"]["skipped"], 3)
        with patch.object(sync_script.subprocess, "run", return_value=CompletedProcess([], 0, '{"rule_failed":1}\n', '')), \
             patch.object(sync_script, "write_state"):
            with self.assertRaises(RuntimeError):
                sync_script.run_script(["parse_grn.py"])

    def test_recent_failed_sync_never_reported_fresh(self):
        from app.services import erp_sync
        with patch.object(erp_sync, "is_running", return_value=False), \
             patch.object(erp_sync, "staleness_hours", return_value=0.1), \
             patch.object(erp_sync, "_read_state", return_value={"state": "failed"}), \
             patch.object(erp_sync, "start", return_value={"state": "running"}) as start:
            self.assertEqual(erp_sync.ensure_fresh()["state"], "running")
            start.assert_called_once()

    def test_zero_attachment_list_age_always_refreshes_and_three_empty_details_pause(self):
        from scripts import erp_fetch_grn_mcp, erp_sync as sync_script
        from unittest.mock import Mock
        self.assertTrue(erp_fetch_grn_mcp.list_is_stale({"fetchedAt": "2999-01-01T00:00:00+08:00"}, 0))
        client = Mock()
        client.call_text.return_value = '{"success":true,"value":{}}'
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(sync_script, "CACHE_ROOT", Path(directory)), \
             patch.object(sync_script, "read_jsonl", return_value=[{"invoiceCode": "missing"+str(i)} for i in range(4)]), \
             patch.object(sync_script, "write_state"), \
             patch("scripts.mcp_erp.McpClient", return_value=client), \
             patch("app.services.feishu_client.get_config", return_value="test"):
            result = sync_script.sync_shipment_details(7, 0)
            self.assertEqual(client.call_text.call_count, 3)
            self.assertTrue(result["interrupted"])
            self.assertEqual(result["failed"], 3)

    def test_sync_default_parses_and_lists_only_is_not_complete(self):
        from scripts import erp_sync as sync_script
        def run(lists_only):
            argv = ["erp_sync.py"] + (["--lists-only"] if lists_only else [])
            with patch("sys.argv", argv), patch.object(sync_script, "write_state") as state, \
                 patch.object(sync_script, "sync_lists"), \
                 patch.object(sync_script, "sync_shipment_details", return_value={"requested": 0, "done": 0, "failed": 0}), \
                 patch.object(sync_script, "run_attachments", return_value={"returncode": 0}), \
                 patch.object(sync_script, "run_script", return_value={"returncode": 0}) as parse, \
                 patch("app.services.shipment_index.build", return_value={}):
                sync_script.main()
                finished = next(call.kwargs for call in state.call_args_list if "fullSyncComplete" in call.kwargs and "finishedAt" in call.kwargs)
                if lists_only:
                    self.assertEqual(finished["state"], "incomplete")
                    self.assertFalse(finished["fullSyncComplete"])
                    parse.assert_not_called()
                else:
                    self.assertEqual(finished["state"], "complete")
                    self.assertTrue(finished["parseComplete"])
                    self.assertEqual([Path(call.args[0][0]).name for call in parse.call_args_list], ["parse_grn.py", "select_grn.py"])
        run(True)
        run(False)

    def test_all_purchase_lists_are_candidates_even_without_grn_flags(self):
        from scripts import erp_fetch_grn_mcp
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(erp_fetch_grn_mcp, "CACHE_ROOT", Path(directory)), \
             patch.object(erp_fetch_grn_mcp, "read_jsonl", return_value=[{"purchase_code": "26MT-03T271-HD", "grnQuantity": "0"}]):
            self.assertEqual(erp_fetch_grn_mcp.candidates()[0]["purchase_code"], "26MT-03T271-HD")

    def test_attachment_candidates_split_joint_codes_but_preserve_add_suffix(self):
        from scripts import erp_fetch_grn_mcp
        purchases = [{"purchase_code": "25MT-05Q552Y&ADD"}, {"purchase_code": "26MT-03T271-HD"}]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(erp_fetch_grn_mcp, "CACHE_ROOT", Path(directory)), \
             patch.object(erp_fetch_grn_mcp, "read_jsonl", return_value=purchases):
            folder = Path(directory) / "details" / "shipments"
            folder.mkdir(parents=True)
            (folder / "sample.json").write_text(json.dumps({"productList": [
                {"采购订单号": "25MT-05Q552Y&ADD"},
                {"采购订单号": "26MT-03T271A,26MT-03T271B"},
                {"采购订单号": "26MT-04T272-HD,26MT-04T273-HD"},
            ]}), encoding="utf-8")
            codes = {row["purchase_code"] for row in erp_fetch_grn_mcp.candidates()}
            self.assertEqual(codes, {"25MT-05Q552Y&ADD", "26MT-03T271-HD", "26MT-04T272-HD", "26MT-04T273-HD"})

    def test_interrupted_sync_requires_manual_resume(self):
        from app.services import erp_sync
        with patch.object(erp_sync, "is_running", return_value=False), \
             patch.object(erp_sync, "staleness_hours", return_value=24), \
             patch.object(erp_sync, "_read_state", return_value={"state": "paused", "interrupted": True}), \
             patch.object(erp_sync, "start") as start:
            self.assertEqual(erp_sync.ensure_fresh()["state"], "paused")
            start.assert_not_called()

    def test_list_page_missing_cannot_be_committed_as_complete(self):
        from scripts.erp_fetch_all import fetch_list
        from app.services.erp_lists import SHIPMENTS
        from unittest.mock import Mock
        session = Mock()
        session.fetch_json.side_effect = [{"root": [], "total": 5}]
        with self.assertRaisesRegex(RuntimeError, "未抓齐"):
            fetch_list(session, SHIPMENTS, pause=0)
        session.fetch_json.side_effect = [{"error": "session expired"}]
        with self.assertRaisesRegex(RuntimeError, "未返回列表"):
            fetch_list(session, SHIPMENTS, pause=0)
        session.fetch_json.side_effect = [{"root": [], "total": 0}]
        self.assertEqual(fetch_list(session, SHIPMENTS, pause=0), [])

    def test_list_boundary_duplicates_are_deduped_without_false_missing_pages(self):
        from scripts.erp_fetch_all import fetch_list
        from app.services.erp_lists import SALE_ORDERS
        from unittest.mock import Mock
        session = Mock()
        session.fetch_json.return_value = {"root": "test", "total": 3}
        with patch("app.services.erp_replay.columns_to_rows", side_effect=[
                [{"orderId": "1"}, {"orderId": "2"}],
                [{"orderId": "2"}, {"orderId": "3"}]]):
            self.assertEqual([row["orderId"] for row in fetch_list(session, SALE_ORDERS, pause=0)], ["1", "2", "3"])

    def test_list_totals_row_does_not_end_pagination_early(self):
        from scripts.erp_fetch_all import fetch_list
        from app.services.erp_lists import PURCHASES
        from unittest.mock import Mock
        session = Mock()
        session.fetch_json.return_value = {"root": "test", "total": 3}
        with patch("app.services.erp_replay.columns_to_rows", side_effect=[
                [{"purchase_id": "1"}, {"purchase_id": "2"}, {"purchase_id": "", "purchase_code": "合计"}],
                [{"purchase_id": "3"}, {"purchase_id": "", "purchase_code": "合计"}]]):
            self.assertEqual([row["purchase_id"] for row in fetch_list(session, PURCHASES, pause=0)], ["1", "2", "3"])
            self.assertEqual(session.fetch_json.call_count, 2)

    def test_new_account_projected_list_preserves_old_relation_fields(self):
        from scripts.erp_sync import merge_rows
        merged, added, updated = merge_rows(
            [{"purchase_id": "2465", "orderCode": "PI-26MT-03T271", "amount": "1", "remark": "old"}],
            [{"purchase_id": "2465", "amount": "544110", "remark": ""}], "purchase_id")
        self.assertEqual(merged[0]["orderCode"], "PI-26MT-03T271")
        self.assertEqual(merged[0]["amount"], "544110")
        self.assertEqual(merged[0]["remark"], "")
        self.assertEqual((added, updated), (0, 1))

    def test_mcp_connection_reset_retries_without_losing_result(self):
        from scripts.mcp_erp import McpClient
        client = McpClient("test")
        with patch.object(client, "_post", side_effect=[urllib.error.URLError("connection reset"),
                {"result": {"content": [{"type": "text", "text": "ok"}]}}]) as request, \
             patch("scripts.mcp_erp.time.sleep"):
            self.assertEqual(client.call_text("shipment.find", {"invoiceCode": "test"}), "ok")
            self.assertEqual(request.call_count, 2)

    def test_mcp_empty_text_is_retried_and_never_faked_as_business_json(self):
        from scripts.mcp_erp import McpClient
        client = McpClient("test")
        empty = {"content": [{"type": "text", "text": ""}]}
        valid = {"content": [{"type": "text", "text": '{"success":true,"value":{"id":"1"}}'}]}
        with patch.object(client, "call", side_effect=[empty, valid]) as request, \
             patch("scripts.mcp_erp.time.sleep"):
            self.assertTrue(json.loads(client.call_text("purchase.find", {}))["success"])
            self.assertEqual(request.call_count, 2)
        with patch.object(client, "call", return_value=empty) as request, \
             patch("scripts.mcp_erp.time.sleep"):
            self.assertEqual(client.call_text("purchase.find", {}), "")
            self.assertEqual(request.call_count, 3)

    def test_lookup_incomplete_input_rejects_before_split_without_fetching_context(self):
        rows = [{"报关单号": "223120260000174064", "合同号_1": "26MT-01A001",
                 "报关品名": "钢管", "总价": "100", "币种": "USD"}]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(workspace, "JOBS_DIR", Path(directory) / "jobs"), \
             patch.object(pipeline, "RUNS_DIR", Path(directory) / "runs"), \
             patch.object(shipment_detail, "shipment_invoices_for_contract", return_value=["PI-26MT-01A001"]), \
             patch.object(shipment_detail, "load_shipment_detail", return_value={"baseInfo": {"出运总金额": "150", "币种": "USD"}, "expenseList": []}), \
             patch.object(lookup, "_fetch", side_effect=AssertionError("must not fetch context")) as related, \
             patch.object(pipeline, "run_round") as split, patch.object(pipeline, "review_rows") as costs:
            workspace._update("incomplete", id="incomplete", kind="lookup", state="queued")
            workspace._run_active("incomplete", "lookup", [], rows, False)
            status = workspace.get_job("incomplete")
            self.assertEqual(status["state"], "failed")
            self.assertEqual(status["failureReason"], "incomplete_input")
            self.assertIn("少 50.00", status["inputErrors"][0]["原因"])
            self.assertFalse((workspace.JOBS_DIR / "incomplete" / "rows.json").exists())
            self.assertEqual(TestClient(app).get("/api/jobs/incomplete/rows").status_code, 404)
            related.assert_not_called()
            split.assert_not_called()
            costs.assert_not_called()

    def test_input_completeness_groups_joint_and_individual_contracts_without_double_counting(self):
        mapping = {"26MT-01A001&002": ["PI-26MT-01A001", "PI-26MT-01A002"],
                   "26MT-01A001": ["PI-26MT-01A001"], "26MT-01A002": ["PI-26MT-01A002"]}
        def detail(invoice):
            return {"baseInfo": {"出运总金额": "100" if invoice.endswith("001") else "50", "币种": "USD"}}
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(shipment_detail, "shipment_invoices_for_contract", side_effect=lambda code: mapping[code]), \
             patch.object(shipment_detail, "load_shipment_detail", side_effect=detail):
            path = Path(directory) / "input.xlsx"
            rows = [{"报关单号": str(223120260000174064+i), "合同号_1": contract,
                     "报关品名": "钢管", "总价": amount, "币种": "USD"}
                    for i, (contract, amount) in enumerate([(list(mapping)[0], "70"), (list(mapping)[1], "50"), (list(mapping)[2], "30")])]
            pipeline.write_declarations(rows, path)
            pipeline.validate_input_completeness(path)
            rows.pop()
            pipeline.write_declarations(rows, path)
            with self.assertRaises(pipeline.BatchCalculationError) as error:
                pipeline.validate_input_completeness(path)
            self.assertIn("少 30.00", str(error.exception))

    def test_missing_erp_details_and_currency_mismatch_never_assume_complete(self):
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(shipment_detail, "shipment_invoices_for_contract", return_value=["PI-26MT-01A001"]):
            path = Path(directory) / "input.xlsx"
            pipeline.write_declarations([{"报关单号": "223120260000174064", "合同号_1": "26MT-01A001",
                                          "报关品名": "钢管", "总价": "100", "币种": "USD"}], path)
            with patch.object(shipment_detail, "load_shipment_detail", return_value=None):
                with self.assertRaises(pipeline.BatchCalculationError) as error:
                    pipeline.validate_input_completeness(path)
                self.assertIn("缺少睿贝出运明细", str(error.exception))
            with patch.object(shipment_detail, "load_shipment_detail", return_value={"baseInfo": {"出运总金额": "100", "币种": "EUR"}}):
                with self.assertRaises(pipeline.BatchCalculationError) as error:
                    pipeline.validate_input_completeness(path)
                self.assertIn("币种", str(error.exception))

    def test_input_completeness_keeps_rounding_tolerance_and_rejects_invalid_amounts(self):
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(shipment_detail, "shipment_invoices_for_contract", return_value=["PI-26MT-01A001"]):
            path = Path(directory) / "input.xlsx"
            row = {"报关单号": "223120260000174064", "合同号_1": "26MT-01A001",
                   "报关品名": "钢管", "总价": "100", "币种": "USD"}
            pipeline.write_declarations([row], path)
            with patch.object(shipment_detail, "load_shipment_detail", return_value={"baseInfo": {"出运总金额": "101.46", "币种": "USD"}}):
                pipeline.validate_input_completeness(path)
            for invalid in ["NaN", "Infinity", "bad"]:
                with patch.object(shipment_detail, "load_shipment_detail", return_value={"baseInfo": {"出运总金额": invalid}}):
                    with self.assertRaises(pipeline.BatchCalculationError) as error:
                        pipeline.validate_input_completeness(path)
                    self.assertIn("出运总金额无效", str(error.exception))
                pipeline.write_declarations([{**row, "总价": invalid}], path)
                with self.assertRaises(pipeline.BatchCalculationError):
                    pipeline.validate_input_completeness(path)
                pipeline.write_declarations([row], path)

    def test_customer_fee_can_be_included_or_omitted_but_currency_still_checked(self):
        row = {"报关单号": "310120260514919201", "合同号_1": "26MT-03T271B",
               "报关品名": "不锈钢管件", "总价": "87991.48", "币种": "USD"}
        detail = {"baseInfo": {"出运总金额": "89991.48", "币种": "USD"},
                  "productList": [{"出运金额": "87991.48"}],
                  "expenseList": [{"金额": "2000", "费用备注": "3.2检验费"}]}
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(shipment_detail, "shipment_invoices_for_contract", return_value=["PI-26MT-03T271B"]), \
             patch.object(shipment_detail, "load_shipment_detail", return_value=detail):
            path = Path(directory) / "input.xlsx"
            for total in ["87991.48", "89991.48"]:
                pipeline.write_declarations([{**row, "总价": total}], path)
                pipeline.validate_input_completeness(path)
            pipeline.write_declarations([{**row, "总价": "43995.74"}], path)
            with self.assertRaises(pipeline.BatchCalculationError):
                pipeline.validate_input_completeness(path)
            # 多笔费用允许只包含其中一笔，不能只支持全含或全不含。
            detail["expenseList"].append({"金额": "300"})
            pipeline.write_declarations([{**row, "总价": "88291.48"}], path)
            pipeline.validate_input_completeness(path)
            pipeline.write_declarations([{**row, "币种": "EUR"}], path)
            with self.assertRaisesRegex(pipeline.BatchCalculationError, "币种"):
                pipeline.validate_input_completeness(path)
            pipeline.write_declarations([row], path)
            detail["expenseList"] = []
            with self.assertRaisesRegex(pipeline.BatchCalculationError, "少 2000.00"):
                pipeline.validate_input_completeness(path)

    def test_run_round_preserves_input_and_never_builds_erp_declarations(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            path.write_bytes(b"unchanged input")
            with patch.object(pipeline, "run_split", return_value={"拆出记录": 1}) as split:
                self.assertEqual(pipeline.run_round(path, Path(directory) / "split"), {"stats": {"拆出记录": 1}, "supplemented": 0})
                split.assert_called_once()
                self.assertEqual(path.read_bytes(), b"unchanged input")

    def test_unclaimed_twin_already_costed_is_allowed_but_unaccounted_goods_reject_batch(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            book = Workbook()
            book.active.append(["合同号_1", "出运金额", "处理", "采购单号"])
            book.active.append(["26MT-01A001", "100", "已按孪生行规则补记成本到 26MT-01A001 的记录", "26MT-01A001-HD"])
            path = folder / "拆单_未认领产品行.xlsx"
            book.save(path)
            pipeline.validate_split_completeness(folder)
            book.active.append(["26MT-01A001", "50", "", "26MT-01A001-HD"])
            book.save(path)
            with self.assertRaises(pipeline.BatchCalculationError) as error:
                pipeline.validate_split_completeness(folder)
            self.assertIn("不会从睿贝生成报关行", str(error.exception))

    def test_real_goods_sequences_are_not_merged_and_null_amount_falls_back_to_total(self):
        common = {"报关单号": "223120260000174064", "合同号_1": "26MT-01A001",
                  "报关品名": "钢管", "报关金额": None, "报关重量": "10"}
        rows = [{**common, "商品序号": "1", "总价": "40"}, {**common, "商品序号": "2", "总价": "60"}]
        merged, count = pipeline.merge_declaration_rows(rows)
        self.assertEqual((len(merged), count), (2, 0))
        merged, count = pipeline.merge_declaration_rows([{**r, "商品序号": ""} for r in rows])
        self.assertEqual(merged[0]["总价"], "100.00")
        self.assertEqual(count, 1)
        records = [{"record_id": str(i), "fields": {**r, "报关金额": r["总价"]}} for i, r in enumerate(rows)]
        self.assertEqual(len(lookup._collapse_records(records)), 2)

    def test_template_preserves_real_goods_sequences_and_customs_fields(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            book = Workbook()
            book.active.append(list(template_import.INPUT_COLUMNS) + ["商品序号", "海关编码"])
            for number, amount in [(1, "40"), (2, "60")]:
                book.active.append(["223120260000174064", "26MT-01A001", "钢管", amount,
                                    "USD", "10", str(number), "7304419000"])
            book.save(path)
            rows, errors = template_import.read_template(path)
            self.assertFalse(errors)
            self.assertEqual([r["商品序号"] for r in rows], ["1", "2"])
            self.assertTrue(all(r["海关编码"] == "7304419000" for r in rows))
            self.assertEqual(len(pipeline.merge_declaration_rows(rows)[0]), 2)

    def test_lookup_contract_only_never_expands_into_other_declarations(self):
        sources = [{"record_id": "extra", "fields": {"报关单号": "223120260000174064",
                    "合同号_1": "26MT-01A001", "报关品名": "钢管", "报关金额": "100"}}]
        result = lookup.enrich_rows([{"合同号_1": "26MT-01A001"}], sources=sources)
        self.assertEqual(len(result["rows"]), 1)
        self.assertFalse(result["rows"][0].get("报关单号"))
        self.assertIn("不再按合同", result["warnings"][0])

    def test_ui_costs_delegate_to_offline_logic_instead_of_loading_saved_results(self):
        rows = [{"采购单号": "26MT-01A001-HD"}]
        with patch("scripts.build_cost_match.allocate_all", return_value=[{"amount": "123"}]) as offline:
            self.assertEqual(review.allocate_costs(rows), [{"amount": "123"}])
            offline.assert_called_once_with(rows, policy="amount")

    def test_sync_worker_accepts_complete_real_context_without_erp_supplements(self):
        rows = [{"报关单号": str(223120260000174064+i), "合同号_1": "26MT-01A001",
                 "报关品名": "钢管", "总价": str(amount), "币种": "USD"}
                for i, amount in enumerate([100, 50])]
        results = [{"报关单号": r["报关单号"], "合同号_1": r["合同号_1"],
                    "采购订单号": "26MT-01A001-HD", "报关金额": r["总价"],
                    "采购金额": r["总价"], "异常类型": "正常"} for r in rows]
        def split(_input, out):
            out.mkdir(parents=True, exist_ok=True)
            book = Workbook()
            book.active.append(["报关单号"])
            book.active.append([rows[0]["报关单号"]])
            book.save(out / "拆单明细_全部.xlsx")
            return {"stats": {"拆出记录": 2}, "supplemented": 0}
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(workspace, "JOBS_DIR", Path(directory) / "jobs"), \
             patch.object(pipeline, "RUNS_DIR", Path(directory) / "runs"), \
             patch.object(shipment_detail, "shipment_invoices_for_contract", return_value=["PI-26MT-01A001"]), \
             patch.object(shipment_detail, "load_shipment_detail", return_value={"baseInfo": {"出运总金额": "150", "币种": "USD"}}), \
             patch.object(pipeline, "run_round", side_effect=split) as run, \
             patch.object(pipeline, "review_rows", return_value=results):
            workspace._update("complete", id="complete", kind="sync", state="queued")
            workspace._run_active("complete", "sync", [], rows, False)
            self.assertEqual(workspace.get_job("complete")["state"], "complete")
            self.assertEqual(workspace.result_rows("complete"), results)
            self.assertTrue(all(r["报关单号"] for r in workspace.result_rows("complete")))
            run.assert_called_once()

    def test_sync_isolates_missing_batch_and_allocates_all_successful_batches_together(self):
        contracts = ["26MT-01A001", "26MT-01A002", "26MT-01A003"]
        inputs = [{"record_id": f"rec{i}", "报关单号": str(223120260000174064+i),
                   "合同号_1": code, "报关品名": "钢管", "总价": "100", "报关金额": "100",
                   "币种": "USD"} for i, code in enumerate(contracts)]
        def detail(invoice):
            return {"baseInfo": {"出运总金额": "150" if invoice.endswith("002") else "100", "币种": "USD"}}
        def split(path, out):
            out.mkdir(parents=True, exist_ok=True)
            rows = workspace._read_sheet(path)
            workspace._write_xlsx([{**r, "采购单号": "shared-po"} for r in rows],
                                  out / "拆单明细_全部.xlsx", ["报关单号", "合同号_1", "采购单号"])
            return {"stats": {}, "supplemented": 0}
        def costs(path):
            return [{**r, "采购金额": "80", "报关品名": "钢管", "报关金额": "100", "异常类型": "正常"}
                    for r in workspace._read_sheet(path)]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(workspace, "JOBS_DIR", Path(directory) / "jobs"), \
             patch.object(pipeline, "RUNS_DIR", Path(directory) / "runs"), \
             patch.object(shipment_detail, "shipment_invoices_for_contract", side_effect=lambda c: ["PI-"+c]), \
             patch.object(shipment_detail, "load_shipment_detail", side_effect=detail), \
             patch.object(pipeline, "run_round", side_effect=split) as run, \
             patch.object(pipeline, "review_rows", side_effect=costs) as review_costs:
            workspace._update("mixed", id="mixed", kind="sync", state="queued")
            (workspace.JOBS_DIR / "mixed" / "source_records.json").write_text(json.dumps(inputs), encoding="utf-8")
            workspace._run_active("mixed", "sync", [], inputs, False)
            status = workspace.get_job("mixed")
            self.assertEqual(status["state"], "complete")
            self.assertEqual(status["batchSummary"]["calculatedBatches"], 2)
            self.assertEqual(status["batchSummary"]["rejectedBatches"], 1)
            rows = workspace.result_rows("mixed")
            self.assertEqual(len(rows), 3)
            self.assertEqual([r["合同号_1"] for r in rows if not r.get("_未核算")], [contracts[0], contracts[2]])
            self.assertIsNone(rows[-1]["采购金额"])
            self.assertIn("少 50.00", rows[-1]["异常明细"])
            self.assertEqual(run.call_count, 2)
            review_costs.assert_called_once()
            self.assertEqual(len(workspace._read_sheet(review_costs.call_args.args[0])), 2)
            proposal = feishu_workflow.plan("mixed")
            blocked = next(g for g in proposal["groups"] if g["sourceId"] == "rec1")
            self.assertEqual(blocked["status"], "blocked")
            book = load_workbook(workspace.download_path("mixed", "exceptions"), read_only=True)
            self.assertEqual(book.active.max_row, 2)
            book.close()

    def test_sync_joint_contract_split_failure_isolated_and_all_rejected_never_costed(self):
        inputs = [{"报关单号": str(223120260000174064+i), "合同号_1": c,
                   "报关品名": "钢管", "总价": "50", "币种": "USD"}
                  for i, c in enumerate(["26MT-01A001", "26MT-01A002"])]
        def split(path, out):
            out.mkdir(parents=True, exist_ok=True)
            workspace._write_xlsx(workspace._read_sheet(path), out / "拆单明细_全部.xlsx", list(pipeline.DECLARATION_COLUMNS))
            workspace._write_xlsx([{"合同号_1": "26MT-01A001", "报关单号": inputs[0]["报关单号"],
                                   "说明": "无法拆单", "类型": "拆单失败"}],
                                  out / "拆单_本地异常.xlsx", ["合同号_1", "报关单号", "说明", "类型"])
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(workspace, "JOBS_DIR", Path(directory) / "jobs"), \
             patch.object(pipeline, "RUNS_DIR", Path(directory) / "runs"), \
             patch.object(shipment_detail, "shipment_invoices_for_contract", return_value=["PI-26MT-01A001&002"]), \
             patch.object(shipment_detail, "load_shipment_detail", return_value={"baseInfo": {"出运总金额": "100", "币种": "USD"}}), \
             patch.object(pipeline, "run_round", side_effect=split), \
             patch.object(pipeline, "review_rows") as cost:
            workspace._update("joint", id="joint", kind="sync", state="queued")
            workspace._run_active("joint", "sync", [], inputs, False)
            self.assertEqual(workspace.get_job("joint")["batchSummary"], {"batches": 1, "calculatedBatches": 0, "rejectedBatches": 1})
            self.assertTrue(all(r["_未核算"] and r["采购金额"] is None for r in workspace.result_rows("joint")))
            cost.assert_not_called()

    def test_partial_pdf_parse_does_not_calculate_surviving_subset(self):
        def parse(_files, path):
            pipeline.write_declarations([{"报关单号": "223120260000174064", "合同号_1": "26MT-01A001",
                                          "报关品名": "钢管", "总价": "100"}], path)
            return {"rows": 1, "failed": [{"文件": "missing.pdf", "原因": "无法解析"}], "warnings": {}}
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(workspace, "JOBS_DIR", Path(directory) / "jobs"), \
             patch.object(pipeline, "RUNS_DIR", Path(directory) / "runs"), \
             patch.object(pipeline, "parse_pdfs", side_effect=parse), \
             patch.object(pipeline, "validate_input_completeness") as validation, \
             patch.object(pipeline, "run_round") as split:
            workspace._update("partial_pdf", id="partial_pdf", kind="files", state="queued")
            workspace._run_active("partial_pdf", "files", [Path(directory)/"input.pdf"], [], False)
            status = workspace.get_job("partial_pdf")
            self.assertEqual(status["state"], "failed")
            self.assertIn("missing.pdf", status["message"])
            self.assertFalse((workspace.JOBS_DIR/"partial_pdf"/"rows.json").exists())
            validation.assert_not_called()
            split.assert_not_called()

    def test_sync_context_includes_outside_range_parts_of_same_erp_shipment(self):
        target = {"record_id": "target", "fields": {"报关单号": "223120260000174064",
                  "合同号_1": "26MT-01A001", "报关品名": "钢管", "报关金额": "100", "采购金额": ""}}
        outside = {"record_id": "outside", "fields": {"报关单号": "223120260000174065",
                   "合同号_1": "26MT-01A002", "报关品名": "钢管", "报关金额": "50", "采购金额": "25"}}
        class FakeClient:
            def iter_search_records(self, *_args):
                yield target
            def search_records(self, _app, _table, query):
                if query["conditions"][0]["field_name"] == "报关单号":
                    return [target]
                self.context_query = query
                return [target, outside]
        fake = FakeClient()
        with patch.object(feishu_workflow, "_target", return_value=(fake, "app", "table")), \
             patch.object(shipment_detail, "shipment_invoices_for_contract", return_value=["PI-26MT-01A001&002"]):
            chosen = feishu_workflow.select_records(limit=1)
        self.assertEqual([r["record_id"] for r in chosen["context"]], ["outside"])
        self.assertTrue(chosen["context"][0]["contextOnly"])
        self.assertIn("26MT-01A002", [r["value"][0] for r in fake.context_query["conditions"]])

    def test_sync_invalid_context_does_not_prevent_other_batches_starting(self):
        target = {"record_id": "target", "fields": {"报关单号": "223120260000174064",
                  "合同号_1": "26MT-01A001", "报关品名": "钢管", "报关金额": "100", "采购金额": ""}}
        invalid = {"record_id": "bad", "fields": {"报关单号": "223120260000174065",
                   "合同号_1": "26MT-01A001", "报关品名": "", "报关金额": "50"}}
        class FakeClient:
            def iter_search_records(self, *_args):
                yield target
            def search_records(self, _app, _table, query):
                return [target] if query["conditions"][0]["field_name"] == "报关单号" else [target, invalid]
        with patch.object(feishu_workflow, "_target", return_value=(FakeClient(), "app", "table")), \
             patch.object(workspace, "start_job", return_value="partial") as start, \
             patch.object(workspace, "_update"):
            result = feishu_workflow.start_scan(limit=1)
            self.assertEqual(result["id"], "partial")
            self.assertTrue(any("无法参与补齐" in row["reason"] for row in result["skipped"]))
            start.assert_called_once()

    def test_manual_adjustment_stays_writable_and_flag_blocks_group(self):
        source = {"record_id": "rec_1", "报关单号": "223120260000174064",
                  "合同号_1": "26MT-01A001", "报关品名": "钢管", "报关金额": "100.00"}
        row = {"报关单号": source["报关单号"], "合同号_1": source["合同号_1"],
               "报关品名": "钢管", "供应商简称": "甲", "报关金额": "100.00",
               "采购金额": "80.00", "异常类型": "正常", "异常明细": ""}
        with tempfile.TemporaryDirectory() as directory, patch.object(workspace, "JOBS_DIR", Path(directory)):
            system_root = Path(directory) / "system_only"
            system_root.mkdir()
            (system_root / "rows.json").write_text(json.dumps([
                {**row, "异常类型": "找不到出运单", "异常明细": "需核对凭证"}],
                ensure_ascii=False), encoding="utf-8")
            workspace._update("system_only", kind="sync", state="complete",
                              summary={"exceptionRows": 1}, createdAt="2026-09-29T00:00:00+00:00")
            self.assertEqual(workspace.all_exceptions(), [])
            root = Path(directory) / "sample"
            root.mkdir()
            (root / "rows.json").write_text(json.dumps([row], ensure_ascii=False), encoding="utf-8")
            (root / "source_records.json").write_text(json.dumps([source], ensure_ascii=False), encoding="utf-8")
            workspace._update("sample", kind="sync", state="complete", createdAt="2026-09-30T00:00:00+00:00")
            self.assertEqual(feishu_workflow.plan("sample")["summary"]["ready"], 1)
            response = TestClient(app).post("/api/jobs/sample/override", json={
                "index": 0, "amount": "82.00", "operator": "财务甲", "reason": "凭证复核"})
            self.assertEqual(response.status_code, 200)
            proposal = feishu_workflow.plan("sample")
            self.assertEqual(proposal["summary"]["ready"], 1)
            self.assertTrue(proposal["groups"][0]["manualAdjusted"])
            # 上报异常仅允许正常记录；独立任务验证阻断与异常汇总。
            (root / "overrides.json").unlink()
            response = TestClient(app).post("/api/jobs/sample/flag-exception", json={
                "index": 0, "operator": "财务乙", "reason": "供应商凭证金额待核实"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(feishu_workflow.plan("sample")["summary"]["ready"], 0)
            self.assertEqual(workspace.all_exceptions()[0]["reason"], "供应商凭证金额待核实")
            self.assertEqual(workspace.get_job("sample")["summary"]["exceptionRows"], 1)
            self.assertEqual(TestClient(app).post("/api/jobs/sample/flag-exception", json={
                "index": 0, "operator": "财务乙", "reason": "重复"}).status_code, 400)
            response = TestClient(app).post("/api/jobs/sample/override", json={
                "index": 0, "amount": "83.00", "operator": "财务乙", "reason": "人工核准"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(workspace.all_exceptions(), [])
            self.assertEqual(feishu_workflow.plan("sample")["summary"]["ready"], 1)
            self.assertEqual(workspace.result_rows("sample")[0]["采购金额"], "83.00")
            (root / "overrides.json").unlink()
            (root / "push_results.json").write_text(json.dumps([
                {"sourceId": "rec_1", "status": "written"}]), encoding="utf-8")
            self.assertEqual(TestClient(app).post("/api/jobs/sample/flag-exception", json={
                "index": 0, "operator": "财务乙", "reason": "写回后上报"}).status_code, 400)
            self.assertEqual(workspace.list_jobs()[0]["writebackReady"], 0)

    def test_flagged_exception_can_resubmit_full_original_batch(self):
        source = {"record_id": "rec_1", "报关单号": "223120260000174064",
                  "合同号_1": "26MT-01A001", "报关品名": "钢管", "报关金额": "100.00"}
        row = {**source, "供应商简称": "甲", "采购金额": "80.00", "异常类型": "正常"}
        with tempfile.TemporaryDirectory() as directory, patch.object(workspace, "JOBS_DIR", Path(directory)):
            root = Path(directory) / "sample"
            root.mkdir()
            for name, value in (("rows.json", [row]), ("source_records.json", [source]),
                                ("input_rows.json", [{**source, "总价": "100.00"}])):
                (root / name).write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
            workspace._update("sample", kind="sync", state="complete", createdAt="2026-09-30T00:00:00+00:00")
            workspace.flag_exception("sample", 0, "财务甲", "系统把供应商匹配错了")
            with patch.object(workspace, "start_job", return_value="new-task") as start:
                response = TestClient(app).post("/api/jobs/sample/resubmit-exception", json={
                    "index": 0, "operator": "技术乙"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["id"], "new-task")
            self.assertEqual(start.call_args.kwargs["declarations"][0]["总价"], "100.00")
            self.assertEqual(start.call_args.kwargs["source_records"], [source])
            self.assertEqual(workspace.all_exceptions()[0]["state"], "resubmitted")
            self.assertEqual(workspace.all_exceptions()[0]["newJobId"], "new-task")
            self.assertEqual(feishu_workflow.plan("sample")["summary"]["ready"], 0)
            self.assertEqual(TestClient(app).post("/api/jobs/sample/resubmit-exception", json={
                "index": 0, "operator": "技术乙"}).status_code, 400)
            self.assertEqual(TestClient(app).post("/api/jobs/sample/override", json={
                "index": 0, "amount": "81.00", "operator": "技术乙"}).status_code, 400)

    def test_sync_selects_uncosted_rows_without_pdf_and_completes_contract_context(self):
        def record(key, declaration, contract, amount, *, cost="", pdf=False, day=1780272000000):
            return {"record_id": key, "fields": {"报关单号": declaration, "合同号_1": contract,
                "报关品名": "钢管", "商品序号": key, "报关金额": amount, "报关重量": "10",
                "币种": "USD", "采购金额": cost, "pdf原件": [{"file_token": "pdf"}] if pdf else [],
                "出口日期": day}}

        target = record("target", "223120260000174064", "26MT-01A001", "100")
        extra = record("extra", "223120260000174065", "26MT-01A001", "50", cost="25")
        already = record("already", "223120260000174066", "26MT-01B001", "60", cost="30")

        class FakeClient:
            def iter_search_records(self, _app, _table, filter):
                self.candidate_filter = filter
                yield target
                yield already  # 防御性本地校验也应排除已核算行

            def search_records(self, _app, _table, filter):
                if filter["conditions"][0]["field_name"] == "报关单号":
                    return [target]
                return [target, extra]

        fake = FakeClient()
        with patch.object(feishu_workflow, "_target", return_value=(fake, "app", "table")):
            result = feishu_workflow.select_records(start_date="2026-01-01", limit=5)
        self.assertEqual([item["record_id"] for item in result["selected"]], ["target"])
        self.assertEqual([item["record_id"] for item in result["context"]], ["extra"])
        self.assertFalse(result["selected"][0]["pdfTokens"])
        self.assertEqual(result["selection"]["supplemental"], 1)
        self.assertIn("isEmpty", [entry["operator"] for entry in fake.candidate_filter["conditions"]])
        self.assertNotIn("pdf原件", [entry["field_name"] for entry in fake.candidate_filter["conditions"]])

    def test_sync_scan_uses_feishu_rows_and_never_downloads_pdf(self):
        source = {"record_id": "source", "报关单号": "223120260000174064",
                  "合同号_1": "26MT-01A001", "报关品名": "钢管", "报关金额": "100",
                  "pdfTokens": ["evidence"]}
        chosen = {"selected": [source], "context": [], "skipped": [], "selection": {"total": 1}}
        with patch.object(feishu_workflow, "select_records", return_value=chosen), \
             patch.object(workspace, "start_job", return_value="job") as start, \
             patch.object(workspace, "_update"):
            result = feishu_workflow.start_scan(limit=1)
        self.assertEqual(result["id"], "job")
        self.assertEqual(start.call_args.kwargs["declarations"][0]["总价"], "100")
        self.assertNotIn("files", start.call_args.kwargs)

    def test_historical_split_merges_amount_and_keeps_original_members(self):
        common = {"报关单号": "223120260000174064", "合同号_1": "26MT-01A001",
                  "报关品名": "钢管", "商品序号": "1", "报关重量": "10", "采购金额": ""}
        first = {"record_id": "first", "fields": {**common, "报关金额": "40",
                 "pdf原件": [{"file_token": "evidence"}]}}
        second = {"record_id": "second", "fields": {**common, "报关金额": "60",
                  "pdf原件": []}}

        class FakeClient:
            def iter_search_records(self, _app, _table, _filter):
                yield first
                yield second

            def search_records(self, _app, _table, _filter):
                return [first, second]

        with patch.object(feishu_workflow, "_target", return_value=(FakeClient(), "app", "table")):
            result = feishu_workflow.select_records(limit=5)
        self.assertEqual(result["selection"]["total"], 1)
        self.assertEqual(result["selection"]["merged"], 1)
        self.assertEqual(result["selected"][0]["报关金额"], "100.00")
        self.assertEqual(len(result["selected"][0]["sourceMembers"]), 2)
        self.assertNotIn("writebackBlocked", result["selected"][0])

    def test_domestic_row_is_labeled_by_contract_not_invalid_declaration(self):
        domestic = {"record_id": "domestic", "fields": {"报关单号": "",
                    "合同号_1": "26MT-01A001", "报关品名": "钢管",
                    "报关金额": "100", "采购金额": ""}}

        class FakeClient:
            def iter_search_records(self, _app, _table, _filter):
                yield domestic

        with patch.object(feishu_workflow, "_target", return_value=(FakeClient(), "app", "table")):
            result = feishu_workflow.select_records(limit=5)
        self.assertEqual(result["selection"]["total"], 0)
        self.assertEqual(result["skipped"][0]["tradeType"], "内销")
        self.assertEqual(result["skipped"][0]["合同号_1"], "26MT-01A001")
        self.assertNotIn("18 位", result["skipped"][0]["reason"])

    def test_sync_endpoints_never_surface_connection_failure_as_500(self):
        with patch.object(feishu_workflow, "select_records", side_effect=urllib.error.URLError("offline")):
            response = TestClient(app).post("/api/sync/preview", data={"limit": 1})
        self.assertEqual(response.status_code, 503)
        with patch.object(feishu_workflow, "start_scan", return_value={"id": None, "selection": {}}) as start:
            response = TestClient(app).post("/api/sync/scan", data={"limit": 1})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("use_ai", start.call_args.kwargs)

    def test_feishu_paste_and_unique_source_enrichment(self):
        self.assertEqual(lookup.parse_paste("报关单号223120260000174064")["rows"][0]["报关单号"],
                         "223120260000174064")
        pasted = lookup.parse_paste(
            "报关单号\t采购订单号\t报关金额\n"
            "223120260000174064\t26MT-01A001-YH\t100.00"
        )
        self.assertEqual(pasted["rows"][0]["报关金额"], "100.00")
        sources = [{"fields": {"报关单号": "223120260000174064", "合同号_1": "26MT-01A001",
                               "报关品名": "钢管", "报关金额": "100.00", "报关重量": 10,
                               "币种": "USD", "pdf原件": [{"file_token": "pdf"}]}}]
        with patch("app.services.erp_cache_index.lookup", return_value={"shipments": [
                {"invoiceCode": "26MT-01A001"}]}):
            enriched = lookup.enrich_rows(pasted["rows"], sources=sources)
        self.assertEqual(enriched["rows"][0]["合同号_1"], "26MT-01A001")
        self.assertEqual(enriched["rows"][0]["报关品名"], "钢管")
        self.assertEqual(lookup.validate_rows(enriched["rows"]), [])

    def test_pasted_row_prefers_contract_with_batch(self):
        """飞书一行常同时有「合同号」和「合同号_1」，要取带批次的那一个。"""
        rows = lookup.parse_paste(
            "224420260014225760 26MT-03R036 26MT-03R036F 不锈钢管件 56 14"
        )["rows"]
        self.assertEqual(rows[0]["合同号_1"], "26MT-03R036F")
        # 同一行的采购订单号带工厂后缀时，基号仍然当合同号
        rows = lookup.parse_paste("26MT-03R036 26MT-03R036-HD")["rows"]
        self.assertEqual(rows[0]["合同号_1"], "26MT-03R036")
        self.assertEqual(rows[0]["采购订单号"], "26MT-03R036-HD")
        # 联合合同号比基号更完整
        rows = lookup.parse_paste("25MT-07F477 25MT-07F477&26MT-07C330")["rows"]
        self.assertEqual(rows[0]["合同号_1"], "25MT-07F477&26MT-07C330")

    def test_amount_claim_stays_inside_order_family(self):
        """按报关金额认领出运单时，只认同一订单家族的出运单。"""
        from scripts.run_split_knapsack import order_family

        self.assertEqual(order_family("26MT-03P315B"), order_family("25MT-03P315-ADD1-A"))
        self.assertNotEqual(order_family("26MT-03P315B"), order_family("25MT-03T614"))
        self.assertEqual(order_family("25MT-03P495Y-ADD1-A"), order_family("25MT-03P495Y-B"))

    def test_contract_batch_missing_in_erp_reports_instead_of_guessing(self):
        """报关合同号写了批次、睿贝却没有这批出运单时，不能拿同订单其它批次顶替。"""
        shipments = [
            {"invoiceCode": "26MT-03R036A", "shipmentId": "1113",
             "orderCode": "PI-26MT-03R036", "purchaseCode": "26MT-03R036-HD"},
            {"invoiceCode": "26MT-03R036E", "shipmentId": "1334",
             "orderCode": "PI-26MT-03R036", "purchaseCode": "26MT-03R036-JX"},
        ]
        with patch("app.services.erp_cache.read_jsonl", return_value=shipments), \
             patch.object(contract_shipments, "read_jsonl", return_value=shipments), \
             patch.object(contract_shipments, "_SHIPMENTS", None):
            self.assertTrue(shipment_detail.names_batch("26MT-03R036F"))
            self.assertFalse(shipment_detail.names_batch("26MT-03R036"))
            # 批次在睿贝里不存在 → 不定位、不兜底
            self.assertEqual(shipment_detail.shipment_invoices_for_contract("26MT-03R036F"), [])
            # 同名出运单 / 没写批次的合同照旧能定位
            self.assertEqual(shipment_detail.shipment_invoices_for_contract("26MT-03R036E"),
                             ["26MT-03R036E"])
            self.assertEqual(set(shipment_detail.shipment_invoices_for_contract("26MT-03R036")),
                             {"26MT-03R036A", "26MT-03R036E"})

    def test_local_split_failure_becomes_finance_readable_row(self):
        """拆单失败的报关行也要出现在结果里，且说明是给财务看的话。"""
        self.assertEqual(review._friendly_local_note("无出运产品行")[0], "找不到出运单")
        self.assertEqual(
            review._friendly_local_note("睿贝里没有这张出运单（26MT-03R036F），该订单只有 A/B/C/D/E 批")[0],
            "找不到出运单",
        )
        self.assertEqual(
            review._friendly_local_note("310120260519759032: 子集和无精确解（容量 73076.83）")[1],
            "报关金额与睿贝出运金额凑不出精确对应，请核对这一单的报关单与出运单。",
        )
        self.assertEqual(
            review._friendly_local_note("费用口径不唯一：小数点调整:add:by_amount、…")[0],
            "拆单失败",
        )

    def test_lookup_collapses_historical_split_rows(self):
        """财务把拆单结果补回表里的多行，粘贴补全时要并回一条报关行。"""
        def record(key, amount, weight):
            return {"record_id": key, "fields": {
                "报关单号": "223120260000174064", "合同号_1": "26MT-01A001",
                "报关品名": "钢管", "报关金额": amount, "报关重量": weight, "币种": "USD"}}

        sources = [record("first", "40", "10"), record("second", "60", "15")]
        # 稀疏粘贴（只给单号 / 合同号）→ 展开成一条合并后的报关行
        parsed = lookup.parse_paste("223120260000174064 26MT-01A001")
        enriched = lookup.enrich_rows(parsed["rows"], sources=sources)
        self.assertEqual(len(enriched["rows"]), 1)
        self.assertEqual(enriched["rows"][0]["报关金额"], "100")
        self.assertEqual(enriched["rows"][0]["报关重量"], "25")
        self.assertIn("并回一条", enriched["rows"][0]["_补全说明"])
        # 粘贴里就是同一报关单的两行拆完单记录 → 同样只留一条，不重复计金额
        parsed = lookup.parse_paste(
            "223120260000174064 26MT-01A001 钢管 40 10\n"
            "223120260000174064 26MT-01A001 钢管 60 15"
        )
        enriched = lookup.enrich_rows(parsed["rows"], sources=sources)
        self.assertEqual(len(enriched["rows"]), 1)
        self.assertEqual(enriched["rows"][0]["报关金额"], "100")

    def test_template_roundtrip(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "模板.xlsx"
            path.write_bytes(template_import.template_bytes())
            rows, errors = template_import.read_template(path)
            self.assertEqual(rows, [])
            self.assertEqual(errors[0]["message"], "模板没有商品行")
            book = load_workbook(path)
            sheet = book.active
            columns = {cell.value: cell.column for cell in sheet[1]}
            values = {"报关单号": "223120260000174064", "合同号_1": "26MT-01A001",
                      "报关品名": "钢管", "报关金额": 100, "币种": "USD", "报关重量": 10}
            for key, value in values.items():
                sheet.cell(2, columns[key], value)
            book.save(path)
            book.close()
            rows, errors = template_import.read_template(path)
            self.assertEqual(errors, [])
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["总价"], "100")

    def test_writeback_plan_preserves_original_amount(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(workspace, "JOBS_DIR", Path(directory)):
            root = Path(directory) / "sample"
            root.mkdir()
            (root / "source_records.json").write_text(json.dumps([{
                "record_id": "rec_1", "报关单号": "123", "合同号_1": "MT",
                "报关品名": "钢管", "报关金额": "100.00"
            }], ensure_ascii=False), encoding="utf-8")
            rows = [
                {"报关单号": "123", "合同号_1": "MT", "报关品名": "钢管", "采购订单号": "PO1", "供应商简称": "甲",
                 "产品类型": "无缝管", "报关金额": "40.00", "采购金额": "20.00", "异常类型": "正常"},
                {"报关单号": "123", "合同号_1": "MT", "报关品名": "钢管", "采购订单号": "PO2", "供应商简称": "乙",
                 "产品类型": "无缝管", "报关金额": "60.00", "采购金额": "30.00", "异常类型": "正常"},
            ]
            (root / "rows.json").write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
            proposal = feishu_workflow.plan("sample")
            self.assertEqual(proposal["summary"]["ready"], 1)
            self.assertEqual(proposal["groups"][0]["resultAmount"], "100.00")
            self.assertEqual(len(proposal["groups"][0]["children"]), 2)
            rows[1]["报关金额"] = "59.99"
            (root / "rows.json").write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
            proposal = feishu_workflow.plan("sample")
            self.assertEqual(proposal["summary"]["blocked"], 1)
            self.assertIn("不守恒", proposal["groups"][0]["reasons"][-1])

    def test_writeback_api_requires_gateway_token(self):
        # 本机（127.0.0.1 / testclient）默认放行写回，只有远端来源才要求网关口令
        remote = TestClient(app, client=("203.0.113.9", 1234))
        response = remote.post("/api/sync/unknown/push", data={"operator": "测试员"})
        self.assertEqual(response.status_code, 403)
        evidence = remote.get("/api/jobs/unknown/erp-evidence/0")
        self.assertEqual(evidence.status_code, 403)
        comparison = remote.get("/api/jobs/unknown/evidence-comparison")
        self.assertEqual(comparison.status_code, 403)
        local = TestClient(app).post("/api/sync/unknown/push", data={"operator": "测试员"})
        self.assertEqual(local.status_code, 400)  # 本机放行，由业务层给出「没有可写回的任务」

    def test_pdf_evidence_compares_input_amount_field(self):
        with tempfile.TemporaryDirectory() as directory:
            def fake_parse(_pdfs, target):
                book = Workbook()
                sheet = book.active
                sheet.append(["报关单号", "合同号_1", "总价"])
                sheet.append(["223120260000174064", "26MT-01A001", 100])
                book.save(target)
                book.close()
                return {"rows": 1}

            with patch.object(workspace.pipeline, "parse_pdfs", side_effect=fake_parse):
                result = workspace._compare_pdf_evidence([{
                    "报关单号": "223120260000174064", "合同号_1": "26MT-01A001",
                    "报关金额": "100.00",
                }], [Path(directory) / "reference.pdf"], Path(directory))
            self.assertEqual(result["rows"][0]["状态"], "结果一致")

    def test_ambiguous_source_stays_unmatched_without_manual_picker(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(workspace, "JOBS_DIR", Path(directory)):
            root = Path(directory) / "sample"
            root.mkdir()
            sources = [{"record_id": key, "报关单号": "123", "报关品名": "钢管", "报关金额": "40.00"}
                       for key in ("rec_a", "rec_b")]
            (root / "source_records.json").write_text(json.dumps(sources, ensure_ascii=False), encoding="utf-8")
            (root / "rows.json").write_text(json.dumps([{"报关单号": "123", "报关品名": "钢管",
                "报关金额": "40.00", "采购金额": "20.00", "异常类型": "正常"}], ensure_ascii=False), encoding="utf-8")
            self.assertEqual(feishu_workflow.plan("sample")["summary"]["unmatched"], 1)
            with self.assertRaisesRegex(ValueError, "自动匹配"):
                feishu_workflow.set_mapping("sample", 0, "rec_a", "财务甲")

    def test_existing_split_rows_match_business_key_before_write(self):
        first = {"record_id": "rec_a", "fields": {"报关单号": "123", "合同号_1": "MT",
                 "供应商简称": "甲", "报关品名": "钢管", "报关金额": 40}}
        second = {"record_id": "rec_b", "fields": {"报关单号": "123", "合同号_1": "MT",
                  "供应商简称": "乙", "报关品名": "钢管", "报关金额": 60}}
        members = [{"record_id": "rec_a", "报关金额": "40"},
                   {"record_id": "rec_b", "报关金额": "60"}]
        children = [{"报关单号": "123", "合同号_1": "MT", "供应商简称": supplier,
                     "报关品名": "钢管"} for supplier in ("乙", "甲")]
        targets = feishu_workflow._write_targets(
            {"sourceId": "rec_a", "originalAmount": "100", "children": children},
            {"sourceMembers": members}, {"rec_a": first, "rec_b": second})
        self.assertEqual(targets, ["rec_b", "rec_a"])

    def test_writeback_preview_blocks_unmatched_legacy_member(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(workspace, "JOBS_DIR", Path(directory)):
            root = Path(directory) / "sample"
            root.mkdir()
            members = [{"record_id": key, "报关单号": "123", "合同号_1": "MT",
                        "报关品名": "钢管", "供应商简称": supplier, "报关金额": amount}
                       for key, supplier, amount in (("a", "甲", "40"), ("b", "乙", "60"))]
            sources = [{"record_id": "a", "报关单号": "123", "合同号_1": "MT",
                        "报关品名": "钢管", "报关金额": "100", "sourceMembers": members}]
            rows = [{"报关单号": "123", "合同号_1": "MT", "报关品名": "钢管",
                     "供应商简称": supplier, "报关金额": amount, "采购金额": "20",
                     "异常类型": "正常"}
                    for supplier, amount in (("甲", "40"), ("乙", "60"))]
            (root / "source_records.json").write_text(json.dumps(sources, ensure_ascii=False), encoding="utf-8")
            (root / "rows.json").write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
            self.assertEqual(feishu_workflow.plan("sample")["summary"]["ready"], 1)
            members[1]["供应商简称"] = "丙"
            (root / "source_records.json").write_text(json.dumps(sources, ensure_ascii=False), encoding="utf-8")
            proposal = feishu_workflow.plan("sample")
            self.assertEqual(proposal["summary"]["blocked"], 1)
            self.assertIn("历史拆分行无法全部", proposal["groups"][0]["reasons"][-1])

    def test_push_reuses_historical_split_rows_by_business_key(self):
        members = [{"record_id": "rec_a", "报关金额": "40"},
                   {"record_id": "rec_b", "报关金额": "60"}]
        live = [
            {"record_id": "rec_a", "fields": {"报关单号": "123", "合同号_1": "MT",
                "供应商简称": "甲", "报关品名": "钢管", "报关金额": 40}},
            {"record_id": "rec_b", "fields": {"报关单号": "123", "合同号_1": "MT",
                "供应商简称": "乙", "报关品名": "钢管", "报关金额": 60}},
        ]
        children = [{"splitKey": key, "报关单号": "123", "合同号_1": "MT",
                     "供应商简称": supplier, "报关品名": "钢管", "报关金额": amount,
                     "采购金额": cost, "报关重量": "5"}
                    for key, supplier, amount, cost in (("b", "乙", "60", "30"),
                                                        ("a", "甲", "40", "20"))]
        proposal = {"summary": {"ready": 1}, "groups": [{"sourceId": "rec_a",
                    "status": "ready", "originalAmount": "100", "children": children}]}

        class FakeClient:
            def __init__(self):
                self.updates = []

            def list_fields(self, _app, _table):
                names = ("报关单号", "合同号_1", "供应商简称", "报关品名", "报关金额",
                         "采购金额", "报关重量", feishu_workflow.META_SOURCE,
                         feishu_workflow.META_KEY, feishu_workflow.META_ORIGINAL)
                numeric = {"报关金额", "采购金额", "报关重量", feishu_workflow.META_ORIGINAL}
                return [{"field_name": name, "type": 2 if name in numeric else 1} for name in names]

            def iter_records(self, _app, _table):
                return iter(live)

            def create_record(self, *_args):
                raise AssertionError("已有联合主键行，不应新增")

            def update_record(self, _app, _table, record_id, values):
                self.updates.append((record_id, values))

        fake = FakeClient()
        with tempfile.TemporaryDirectory() as directory, patch.object(workspace, "JOBS_DIR", Path(directory)), \
             patch.object(feishu_workflow, "_target", return_value=(fake, "app", "table")), \
             patch.object(feishu_workflow, "plan", return_value=proposal), \
             patch.object(workspace, "append_audit"):
            root = Path(directory) / "sample"
            root.mkdir()
            (root / "source_records.json").write_text(json.dumps([{
                "record_id": "rec_a", "sourceMembers": members}], ensure_ascii=False), encoding="utf-8")
            result = feishu_workflow.push("sample", "财务甲")
        self.assertEqual((result["written"], result["failed"]), (1, 0))
        self.assertEqual(result["results"][0]["created"], 0)
        self.assertEqual([(record_id, values["报关金额"]) for record_id, values in fake.updates],
                         [("rec_b", 60.0), ("rec_a", 40.0)])

    def test_push_updates_original_with_first_split_amount(self):
        class FakeClient:
            def __init__(self):
                self.calls = []

            def list_fields(self, _app, _table):
                names = ("报关单号", "合同号_1", "供应商简称", "报关品名", "报关金额", "采购金额", "报关重量",
                         feishu_workflow.META_SOURCE, feishu_workflow.META_KEY,
                         feishu_workflow.META_ORIGINAL)
                return [{"field_name": name, "type": 2 if name in ("报关金额", "采购金额", "报关重量", feishu_workflow.META_ORIGINAL) else 1}
                        for name in names]

            def iter_records(self, _app, _table):
                return iter([{"record_id": "rec_1", "fields": {"报关金额": 100.0}}])

            def create_record(self, _app, _table, values):
                self.calls.append(("create", values))
                return {"record": {"record_id": "rec_child"}}

            def update_record(self, _app, _table, record_id, values):
                self.calls.append(("update", record_id, values))

        children = [
            {"splitKey": "first", "报关单号": "123", "合同号_1": "MT", "供应商简称": "甲", "报关品名": "钢管",
             "报关金额": "40.00", "采购金额": "20.00", "报关重量": "4"},
            {"splitKey": "second", "报关单号": "123", "合同号_1": "MT", "供应商简称": "乙", "报关品名": "钢管",
             "报关金额": "60.00", "采购金额": "30.00", "报关重量": "6"},
        ]
        proposal = {"summary": {"ready": 1}, "groups": [{"sourceId": "rec_1", "status": "ready",
                    "originalAmount": "100.00", "children": children}], "unmatched": []}
        fake = FakeClient()
        with tempfile.TemporaryDirectory() as directory, patch.object(workspace, "JOBS_DIR", Path(directory)), \
             patch.object(feishu_workflow, "_target", return_value=(fake, "app", "table")), \
             patch.object(feishu_workflow, "plan", return_value=proposal), \
             patch.object(workspace, "append_audit"):
            (Path(directory) / "sample").mkdir()
            result = feishu_workflow.push("sample", "测试员")
        self.assertEqual(result["written"], 1)
        self.assertEqual(fake.calls[0][0], "create")
        self.assertEqual(fake.calls[0][1]["报关金额"], 60.0)
        self.assertEqual(fake.calls[1][0:2], ("update", "rec_1"))
        self.assertEqual(fake.calls[1][2]["报关金额"], 40.0)
        self.assertEqual(fake.calls[1][2]["采购金额"], 20.0)

        class BrokenClient(FakeClient):
            def create_record(self, _app, _table, values):
                raise RuntimeError("模拟新增子行失败")

        broken = BrokenClient()
        with tempfile.TemporaryDirectory() as directory, patch.object(workspace, "JOBS_DIR", Path(directory)), \
             patch.object(feishu_workflow, "_target", return_value=(broken, "app", "table")), \
             patch.object(feishu_workflow, "plan", return_value=proposal), \
             patch.object(workspace, "append_audit"):
            (Path(directory) / "sample").mkdir()
            result = feishu_workflow.push("sample", "测试员")
        self.assertEqual(result["failed"], 1)
        self.assertFalse(any(call[0] == "update" and call[1] == "rec_1" for call in broken.calls))


if __name__ == "__main__":
    unittest.main()
