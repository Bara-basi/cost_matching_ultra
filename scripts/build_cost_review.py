"""把「与飞书/睿贝不一致」的单元整理成一张带证据的待确认表。

每条单元给出：

* 我方金额与**依据**（入库单哪个批次区块 / 出运采购金额直接取值 / 按人民币占比分摊）
* 飞书金额、睿贝采购金额
* 该采购单的有效入库单及其各批次区块金额（这是我们的原始依据）
* 判定：我方是否有独立证据（入库单区块 + 睿贝出运/采购金额）

输出：`outputs/cost_match/待确认清单.xlsx`
"""
from __future__ import annotations

import collections
import json
import sys
from decimal import Decimal
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.services.cost_match import dec  # noqa: E402
from app.services.grn_extract import PARSED_ROOT, safe_name  # noqa: E402
from app.services.grn_select import erp_amount, select  # noqa: E402

OUT_DIR = PROJECT_ROOT / "outputs" / "cost_match"
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFE699")


def read_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(cell) for cell in rows[0]]
    out = [dict(zip(header, row)) for row in rows[1:]]
    workbook.close()
    return out


def main() -> None:
    units = read_rows(OUT_DIR / "采购金额_单元差异.xlsx")
    details = read_rows(OUT_DIR / "采购金额_匹配.xlsx")
    # 已人工核实过的冲突（结论与证据都写在 data/reference/verified_conflicts.json）
    verified: dict[str, dict] = {}
    verified_path = PROJECT_ROOT / "data" / "reference" / "verified_conflicts.json"
    if verified_path.exists():
        payload = json.loads(verified_path.read_text(encoding="utf-8"))
        for case in payload.get("cases") or []:
            decl = str(case.get("报关单号") or "").strip()
            if decl:
                verified[decl] = case
    by_unit: dict[tuple[str, str, str], list[dict]] = collections.defaultdict(list)
    for row in details:
        key = (str(row["报关单号"]), str(row["供应商简称"]), str(row["采购单核心"]))
        by_unit[key].append(row)

    evidence_cache: dict[str, dict] = {}

    def evidence(code: str) -> dict:
        if code in evidence_cache:
            return evidence_cache[code]
        result = select(code)
        blocks: list[str] = []
        files: list[str] = []
        paths: list[str] = []
        fee_total = Decimal(0)
        for item in result["kept"]:
            files.append(str(item["original"]))
            paths.append(f".cache/erp/attachments/grn/{code}/{item['file']}")
            for block in item.get("blocks") or []:
                if block.get("amount"):
                    blocks.append(f"{block.get('label')}={block.get('amount')}")
            payload_path = PARSED_ROOT / safe_name(code) / (
                safe_name(item["file"], "file") + ".json"
            )
            if payload_path.exists():
                payload = json.loads(payload_path.read_text(encoding="utf-8"))
                # fee_by_column 与 extra_fees 是同一批费用行的两种写法，只能取一种，
                # 否则同一笔费用会被算两遍
                by_column = payload.get("fee_by_column") or {}
                if by_column:
                    fee_total += sum((dec(v) for v in by_column.values()), Decimal(0))
                else:
                    for fee in payload.get("extra_fees") or []:
                        if isinstance(fee, dict) and fee.get("amount"):
                            fee_total += dec(fee.get("amount"))
        info = {
            "files": files,
            "paths": paths,
            "blocks": blocks,
            "dropped": [str(item["original"]) for item in result["dropped"]],
            "issues": list(result["issues"]),
            "fee_total": fee_total,
        }
        evidence_cache[code] = info
        return info

    rows: list[list] = []
    for unit in units:
        verdict = str(unit.get("结论") or "")
        if verdict.startswith("一致"):
            continue
        decl = str(unit["报关单号"])
        vendor = str(unit["供应商简称"])
        core = str(unit["采购单核心"])
        mine = dec(unit["采购金额_我方"])
        theirs = unit["采购金额_飞书"]
        erp = dec(unit["ERP采购金额"])
        detail_rows = by_unit.get((decl, vendor, core), [])
        basis = "；".join(
            f"{row['采购单号']}：{row['采购金额_我方']}（{row['分摊依据']}）"
            for row in detail_rows
        )
        codes = [str(row["采购单号"]) for row in detail_rows]
        blocks: list[str] = []
        files: list[str] = []
        paths: list[str] = []
        fee_total = Decimal(0)
        dropped: list[str] = []
        for code in codes:
            info = evidence(code)
            blocks.extend(f"[{code}] {b}" for b in info["blocks"])
            files.extend(f"[{code}] {f}" for f in info["files"])
            paths.extend(info["paths"])
            dropped.extend(f"[{code}] {d}" for d in info["dropped"])
            fee_total += info.get("fee_total") or Decimal(0)
        erp_sum = sum((erp_amount(code) or Decimal(0) for code in codes), Decimal(0))
        # 判定：我方金额的依据来自哪一层证据
        support: list[str] = []
        for row in detail_rows:
            basis_text = str(row.get("分摊依据") or "")
            if basis_text.startswith("入库单区块"):
                tag = "入库单批次区块直取"
            elif "直接取值" in basis_text:
                tag = "出运采购金额直接取值（只覆盖采购单一部分）"
            elif basis_text.startswith("本记录出运采购金额"):
                tag = "按本记录出运采购金额占比分摊"
            elif basis_text.startswith("ERP出运采购金额"):
                tag = "按睿贝出运采购金额（按产品类型）占比分摊"
            elif basis_text.startswith("出运金额"):
                tag = "按出运金额(USD)占比分摊"
            elif "均分" in basis_text:
                tag = "无权重，均分"
            elif "兜底" in basis_text or "ERP" in basis_text:
                tag = "按睿贝采购金额兜底"
            else:
                tag = basis_text[:20]
            if tag and tag not in support:
                support.append(tag)
        if erp and abs(mine - erp) / erp <= Decimal("0.01"):
            support.append("单元合计等于睿贝采购金额")
        # 给出建议：已核实过的直接引用结论；其余按证据强度
        case = verified.get(decl)
        if case:
            suggestion = f"已核实：{case.get('结论')}（{case.get('核实日期')}）"
        elif (
            erp
            and fee_total > 0
            and abs((mine - erp) - fee_total) <= Decimal("0.05")
        ):
            suggestion = (
                f"建议以我方为准（差额 {fee_total} 正是单据内费用行：包装费/木箱费等；"
                "睿贝与飞书用的是下单金额）"
            )
        elif "入库单批次区块直取" in "、".join(support):
            suggestion = "建议以我方为准（依据 = 入库单批次区块）"
        elif "单元合计等于睿贝采购金额" in "、".join(support):
            suggestion = "建议以我方为准（与睿贝采购金额一致）"
        elif "无法分摊" in str(unit.get("结论")) or "未定价" in "、".join(support):
            suggestion = "需人工：睿贝明细未定价"
        elif theirs not in (None, ""):
            try:
                if abs(mine - dec(theirs)) / dec(theirs) <= Decimal("0.03"):
                    suggestion = "接近飞书（≤3%），建议人工确认口径"
                else:
                    suggestion = "需人工核对口径与依据"
            except Exception:  # noqa: BLE001
                suggestion = "需人工核对口径与依据"
        else:
            suggestion = "需人工核对口径与依据"
        rows.append(
            [
                decl,
                vendor,
                core,
                "、".join(codes),
                str(mine),
                "" if theirs in (None, "") else str(theirs),
                str(erp_sum) if erp_sum else "",
                unit.get("偏差_飞书") or "",
                unit.get("偏差_ERP") or "",
                verdict,
                "、".join(support) if support else "需人工判断",
                suggestion,
                basis,
                "；".join(blocks),
                "；".join(files),
                "；".join(paths),
                "；".join(dropped),
            ]
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "待确认"
    header = [
        "报关单号", "供应商简称", "采购单核心", "采购单号", "我方采购金额", "飞书采购金额",
        "睿贝采购金额", "偏差_飞书", "偏差_ERP", "系统判定", "我方证据", "建议", "分摊依据",
        "入库单实发区块", "有效入库单", "附件路径", "已丢弃附件",
    ]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
        fill = WARN_FILL if "需人工" in str(row[10]) else None
        if fill:
            for cell in sheet[sheet.max_row]:
                cell.fill = fill
    for index, name in enumerate(header, 1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            12, min(60, len(name) * 2 + 4)
        )
    sheet.freeze_panes = "A2"
    path = OUT_DIR / "待确认清单.xlsx"
    workbook.save(path)
    print(f"written {path} ({len(rows)} rows)")


if __name__ == "__main__":
    main()
