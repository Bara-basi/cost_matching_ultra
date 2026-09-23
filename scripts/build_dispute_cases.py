"""把「差额较大」的单元整理成逐案对账表（附原表关键数值）。

一行 = 一个单元，列出：我方金额与算式、飞书原行、睿贝采购金额、
入库单原件里的订单金额/实发金额/各批次区块/费用行/备注，以及拆单记录的金额口径。
用于人工逐条核对"到底哪一侧的数据有问题"。

输出：`outputs/cost_match/大额差异_原表对账.xlsx`
"""
from __future__ import annotations

import json
import sys
from decimal import Decimal
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.services.contract_shipments import cores_of  # noqa: E402
from app.services.cost_match import dec, unit_core  # noqa: E402
from app.services.grn_extract import PARSED_ROOT, safe_name  # noqa: E402
from app.services.grn_select import erp_amount, erp_supplier, select  # noqa: E402

OUT_DIR = PROJECT_ROOT / "outputs" / "cost_match"
REF_CACHE = PROJECT_ROOT / "data" / "cache" / "ref_2026_customs.json"
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")


def read_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(cell) for cell in rows[0]]
    out = [dict(zip(header, row)) for row in rows[1:]]
    workbook.close()
    return out


def po_core(code: str) -> str:
    return unit_core(code)


def main() -> None:
    units = read_rows(OUT_DIR / "采购金额_单元差异.xlsx")
    details = read_rows(OUT_DIR / "采购金额_匹配.xlsx")
    by_unit: dict[tuple, list[dict]] = {}
    for row in details:
        key = (row["报关单号"], row["供应商简称"], row["采购单核心"])
        by_unit.setdefault(key, []).append(row)

    ref = json.loads(REF_CACHE.read_text(encoding="utf-8"))["records"]
    feishu: dict[tuple, list[dict]] = {}
    for row in ref:
        key = (
            str(row.get("报关单号") or "").strip(),
            str(row.get("供应商简称") or "").strip(),
            po_core(str(row.get("采购单号") or "")),
        )
        feishu.setdefault(key, []).append(row)

    def big(unit: dict) -> bool:
        for field, other in (("偏差_飞书", unit.get("采购金额_飞书")), ("偏差_ERP", unit.get("ERP采购金额"))):
            text = str(unit.get(field) or "").replace("%", "")
            if not text or other in (None, ""):
                continue
            try:
                if Decimal(text) > Decimal("10"):
                    return True
            except Exception:  # noqa: BLE001
                continue
        return False

    rows: list[list] = []
    for unit in units:
        if str(unit.get("结论", "")).startswith("一致"):
            continue
        if not big(unit):
            continue
        key = (unit["报关单号"], unit["供应商简称"], unit["采购单核心"])
        record_rows = by_unit.get(key, [])
        codes = [str(r["采购单号"]) for r in record_rows]
        basis = "；".join(
            f"{r['采购单号']}：{r['采购金额_我方']}（{r['分摊依据'][:40]}）" for r in record_rows
        )
        shipment = "；".join(
            f"{r['采购单号']} 报关金额 {r['出运金额合计(USD)']} USD / 出运采购金额 {r.get('出运采购金额合计', '') if '出运采购金额合计' in r else ''}"
            for r in record_rows
        )
        feishu_text = "；".join(
            f"{r.get('报关品名')} 采购金额 {r.get('采购金额')}（报关金额 {r.get('报关金额')}，重量 {r.get('报关重量')}）"
            for r in feishu.get(key, [])
        )
        # 入库单原件关键数值
        book_parts: list[str] = []
        for code in codes:
            result = select(code)
            for item in result["kept"]:
                payload_path = PARSED_ROOT / safe_name(code) / (
                    safe_name(item["file"], "file") + ".json"
                )
                payload = {}
                if payload_path.exists():
                    payload = json.loads(payload_path.read_text(encoding="utf-8"))
                blocks = "、".join(
                    f"{b.get('label')}={b.get('amount')}" for b in (payload.get("settled_blocks") or [])
                )
                fees = "、".join(
                    f"{f.get('label')}={f.get('amount')}" for f in (payload.get("extra_fees") or [])
                )
                book_parts.append(
                    f"[{code}] {item['original']}：订单金额={payload.get('order_amount')}，"
                    f"实发金额={item['amount']}，区块=[{blocks}]，费用=[{fees}]，备注={str(payload.get('notes'))[:60]}"
                )
        erp_text = "；".join(
            f"{code}：{erp_amount(code)}（{erp_supplier(code)[:16]}）" for code in codes
        )
        rows.append(
            [
                unit["报关单号"],
                unit["供应商简称"],
                unit["采购单核心"],
                "、".join(codes),
                unit["采购金额_我方"],
                unit.get("采购金额_飞书"),
                unit.get("ERP采购金额"),
                unit.get("偏差_飞书"),
                unit.get("偏差_ERP"),
                unit.get("结论"),
                basis,
                shipment,
                feishu_text,
                "；".join(book_parts),
                erp_text,
            ]
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "大额差异对账"
    header = [
        "报关单号", "供应商简称", "采购单核心", "采购单号", "我方采购金额", "飞书采购金额",
        "睿贝采购金额", "偏差_飞书", "偏差_ERP", "系统判定", "我方算式", "拆单记录口径",
        "飞书原行", "入库单原件关键数值", "睿贝采购单",
    ]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.fill = BAD_FILL
    for index, name in enumerate(header, 1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            14, min(70, len(name) * 2 + 6)
        )
    sheet.freeze_panes = "A2"
    path = OUT_DIR / "大额差异_原表对账.xlsx"
    workbook.save(path)
    print(f"written {path} ({len(rows)} rows)")


if __name__ == "__main__":
    main()
