"""把 PDF 解析结果与飞书「2026年报关记录 AI」按报关单号比对，评估还原度。

比对口径：
- 报关单号 / 合同号_1 / 报关品名 / 海关编码：逐份报关单直接比对；
- 报关重量：按报关单号汇总后比对（飞书会按产品类型/供应商再拆行）。
"""
from __future__ import annotations

import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
CACHE = PROJECT_ROOT / "data" / "cache"
XLSX = PROJECT_ROOT / "outputs" / "customs_parse" / "报关单解析结果_出口退税联.xlsx"


def flatten(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return " | ".join(filter(None, (flatten(v) for v in value)))
    if isinstance(value, dict):
        for key in ("text", "name", "value"):
            if key in value:
                return flatten(value[key])
    return ""


def norm(value) -> str:
    return "".join(str(value or "").split()).upper()


def as_number(value) -> float | None:
    text = flatten(value).replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def main() -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(XLSX, read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    header = [str(x) for x in rows[0]]
    parsed = [dict(zip(header, row)) for row in rows[1:]]

    records = json.loads((CACHE / "records_ai.json").read_text(encoding="utf-8"))
    by_decl: dict[str, list[dict]] = defaultdict(list)
    for record in records:
        fields = record.get("fields", {})
        decl = norm(flatten(fields.get("报关单号")))
        if decl:
            by_decl[decl].append(fields)

    # PDF 侧按报关单号汇总
    pdf: dict[str, dict] = {}
    dup_files: dict[str, set[str]] = defaultdict(set)
    for row in parsed:
        decl = norm(row.get("报关单号"))
        if not decl:
            continue
        entry = pdf.setdefault(
            decl,
            {
                "contracts": set(),
                "names": set(),
                "hs": set(),
                "weight": 0.0,
                "files": set(),
            },
        )
        entry["contracts"].add(norm(row.get("合同号_1")))
        entry["names"].add(norm(row.get("报关品名")))
        entry["hs"].add(norm(row.get("海关编码")))
        entry["files"].add(str(row.get("来源文件")))
        weight = as_number(row.get("报关重量"))
        if weight is not None:
            entry["weight"] += weight
    for decl, entry in pdf.items():
        if len(entry["files"]) > 1:
            dup_files[decl] = entry["files"]

    stats = Counter()
    detail: list[str] = []
    for decl, entry in pdf.items():
        stats["PDF报关单数"] += 1
        # 同一报关单在 PDF 里可能有多个副本文件，重量只按一份计
        copy_count = max(1, len(entry["files"]))
        entry["weight_single"] = entry["weight"] / copy_count
        matches = by_decl.get(decl)
        if not matches:
            stats["飞书无此报关单"] += 1
            detail.append(f"[无对应] {decl} {sorted(entry['contracts'])}")
            continue
        stats["飞书有此报关单"] += 1
        if entry["contracts"] & {norm(flatten(f.get("合同号_1"))) for f in matches}:
            stats["合同号_1一致"] += 1
        else:
            detail.append(f"[合同号不一致] {decl} PDF={sorted(entry['contracts'])}")
        names = {norm(flatten(f.get("报关品名"))) for f in matches}
        hs = {norm(flatten(f.get("海关编码"))) for f in matches}
        if entry["names"] == names:
            stats["报关品名一致"] += 1
        else:
            stats["报关品名不一致"] += 1
            detail.append(f"[品名不一致] {decl} PDF={sorted(entry['names'])} 飞书={sorted(names)}")
        if entry["hs"] == hs:
            stats["海关编码一致"] += 1
        else:
            stats["海关编码不一致"] += 1
            detail.append(f"[编码不一致] {decl} PDF={sorted(entry['hs'])} 飞书={sorted(hs)}")

        table_weight = sum(
            (as_number(f.get("报关重量")) or 0.0) for f in matches
        )
        if abs(table_weight - entry["weight_single"]) <= 0.05:
            stats["重量合计一致"] += 1
        else:
            stats["重量合计不一致"] += 1
            detail.append(
                f"[重量不一致] {decl} PDF={entry['weight_single']} 飞书={round(table_weight, 2)}"
            )

    out = [f"{k}={v}" for k, v in stats.items()]
    out.append(f"PDF 中同一报关单多份文件的组数={len(dup_files)}")
    out.append("")
    out.extend(detail[:40])
    target = PROJECT_ROOT / "tmp" / "verify_parse.txt"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("\n".join(out), encoding="utf-8")
    print(" | ".join(f"{k}={v}" for k, v in stats.items()))


if __name__ == "__main__":
    main()
