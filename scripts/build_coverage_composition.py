"""飞书有采购金额的行，按「能不能按报关单匹配」分类，看清每一块钱卡在哪。

分类：
  1. 有报关单号 + MT 合同号 → 已匹配；
  2. 有报关单号 + MT 合同号 → 未匹配（缺报关单 PDF，用入库单余额已印证大部分）；
  3. 没有报关单号（内销/视同销售）→ 本来就不是出口报关行，不在按报关单匹配的范围；
  4. 供应商自家编号（CY-/SM- 系列）→ 不是 MT 采购单。

输出：outputs/cost_match/覆盖构成.xlsx
"""
from __future__ import annotations

import json
import re
import sys
from decimal import Decimal
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.services.cost_match import unit_core  # noqa: E402
from app.services.erp_cache import CACHE_ROOT, read_jsonl  # noqa: E402
from app.services.grn_extract import PARSED_ROOT, safe_name  # noqa: E402
from app.services.grn_select import select  # noqa: E402
from app.services.supplier_names import SupplierNames  # noqa: E402

REF = PROJECT_ROOT / "data" / "cache" / "ref_2026_customs_full.json"
MATCH = PROJECT_ROOT / "outputs" / "cost_match" / "采购金额_匹配.xlsx"
OUT = PROJECT_ROOT / "outputs" / "cost_match" / "覆盖构成.xlsx"
UNCOVERED = PROJECT_ROOT / "outputs" / "cost_match" / "缺单行金额核对.xlsx"
FIELDS = [
    "合同号（应收表格）",
    "供应商简称",
    "报关单号",
    "报关品名",
    "报关金额",
    "报关重量",
    "采购金额",
    "开票金额",
    "销售方式",
    "备注",
]
FILLS = {
    "已匹配（报关单号 + MT 合同号）": PatternFill("solid", fgColor="E2EFDA"),
    "未匹配：缺报关单（报关单号 + MT 合同号）": PatternFill("solid", fgColor="FCE4D6"),
    "非报关行：没有报关单号（内销/视同销售）": PatternFill("solid", fgColor="EDEDED"),
    "非 MT 单：供应商自家编号（CY-/SM-）": PatternFill("solid", fgColor="FFF2CC"),
}


def dec(value) -> Decimal:
    try:
        return Decimal(str(value or 0).replace(",", ""))
    except Exception:
        return Decimal(0)


def audit_core(code: str) -> str:
    """比对用的订单核心：`unit_core` 之后再剥掉尾部工厂后缀。

    `unit_core` 对标准单号（26MT-02N182-JX）已剥掉工厂后缀，但对 `26MT-DP002-XMLS`
    这类非标准单号不剥——飞书那边写的是 `26MT-DP002`，不补一刀就永远配不上。
    """
    return re.sub(r"-[A-Z]{2,5}$", "", unit_core(code))


def classify(record: dict) -> str:
    decl = str(record.get("报关单号") or "").strip()
    contract = str(record.get("合同号（应收表格）") or record.get("合同号_1") or "")
    if re.search(r"SM-|SM\d|^CY-", contract.upper()):
        return "非 MT 单：供应商自家编号（CY-/SM-）"
    if not decl:
        return "非报关行：没有报关单号（内销/视同销售）"
    return "已匹配（报关单号 + MT 合同号）"


def po_totals_and_blocks() -> tuple[dict, dict]:
    """我方（订单核心 × 供应商）→ 入库单保留合计 / 各实发区块金额。"""
    names = SupplierNames()

    def norm(text: str) -> str:
        return str(text or "").replace(" ", "").replace("_", "").upper()

    purchases = {
        norm(str(row.get("purchase_code") or "")): names.short(str(row.get("supplierName") or ""))
        for row in read_jsonl(CACHE_ROOT / "purchases" / "purchases.jsonl")
    }
    totals: dict[tuple[str, str], Decimal] = {}
    blocks: dict[tuple[str, str], list[Decimal]] = {}
    for folder in sorted(p for p in PARSED_ROOT.iterdir() if p.is_dir()):
        code = folder.name
        result = select(code)
        if not result["kept"]:
            continue
        key = (audit_core(code), purchases.get(norm(code), ""))
        totals[key] = totals.get(key, Decimal(0)) + sum(
            (dec(item.get("amount")) for item in result["kept"]), Decimal(0)
        )
        table = blocks.setdefault(key, [])
        for item in result["kept"]:
            path = PARSED_ROOT / safe_name(code) / (safe_name(item["file"], "file") + ".json")
            if not path.exists():
                continue
            try:
                payload = json.loads(path.read_text(encoding="utf-8"))
            except ValueError:
                continue
            for block in payload.get("settled_blocks") or []:
                if isinstance(block, dict) and block.get("amount"):
                    table.append(dec(block.get("amount")))
    return totals, blocks


def main() -> None:
    records = [
        record
        for record in json.loads(REF.read_text(encoding="utf-8")).get("records") or []
        if record.get("采购金额")
    ]
    book = load_workbook(MATCH, read_only=True, data_only=True)
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(c) for c in rows[0]]
    covered: set[tuple[str, str]] = set()
    for row in rows[1:]:
        data = dict(zip(header, row))
        covered.add(
            (
                str(data.get("报关单号") or ""),
                audit_core(data.get("采购单核心") or data.get("采购单号")),
            )
        )
    book.close()

    buckets: dict[str, dict] = {}
    detail: list[dict] = []
    for record in records:
        kind = classify(record)
        if kind == "已匹配（报关单号 + MT 合同号）":
            decl = str(record.get("报关单号") or "").strip()
            contract = str(record.get("合同号（应收表格）") or record.get("合同号_1") or "")
            if (decl, audit_core(contract)) not in covered:
                kind = "未匹配：缺报关单（报关单号 + MT 合同号）"
        bucket = buckets.setdefault(kind, {"rows": 0, "money": Decimal(0)})
        bucket["rows"] += 1
        bucket["money"] += dec(record.get("采购金额"))
        detail.append({**{f: record.get(f) for f in FIELDS}, "_kind": kind})

    total = sum((b["money"] for b in buckets.values()), Decimal(0))

    # ---- 印证情况：缺单行看「缺单行金额核对」，内销行按采购单合计/实发区块比
    corroborated: dict[str, Decimal] = {}
    if UNCOVERED.exists():
        ledger = load_workbook(UNCOVERED, read_only=True, data_only=True)
        sheet_ = ledger["缺单行核对"]
        money = Decimal(0)
        for row in list(sheet_.iter_rows(values_only=True))[1:]:
            if str(row[9]).startswith("可采信"):
                money += dec(row[5])
        ledger.close()
        corroborated["未匹配：缺报关单（报关单号 + MT 合同号）"] = money
    totals, blocks = po_totals_and_blocks()
    internal: list[dict] = [r for r in detail if r["_kind"].startswith("非报关行")]
    groups: dict[tuple[str, str], list[dict]] = {}
    for row in internal:
        contract = str(row.get("合同号（应收表格）") or "")
        groups.setdefault((audit_core(contract), str(row.get("供应商简称") or "").strip()), []).append(row)
    internal_ok = Decimal(0)
    for key, members in groups.items():
        money = sum((dec(r.get("采购金额")) for r in members), Decimal(0))
        tol = max(Decimal(1), abs(money) * Decimal("0.01"))
        total_po = totals.get(key)
        if (total_po is not None and abs(total_po - money) <= tol) or any(
            abs(value - money) <= tol for value in blocks.get(key, [])
        ):
            internal_ok += money
    corroborated["非报关行：没有报关单号（内销/视同销售）"] = internal_ok
    corroborated["已匹配（报关单号 + MT 合同号）"] = buckets.get(
        "已匹配（报关单号 + MT 合同号）", {"money": Decimal(0)}
    )["money"]

    workbook = Workbook()
    summary = workbook.active
    summary.title = "构成"
    summary.append(["分类", "行数", "采购金额", "占比", "其中已印证", "说明"])
    notes = {
        "已匹配（报关单号 + MT 合同号）": "我方拆单 + 入库单实发已经算出金额",
        "未匹配：缺报关单（报关单号 + MT 合同号）": "缺的是报关单 PDF；其中绝大部分已用「入库单余额 = 飞书缺单行」印证（见 缺单行金额核对.xlsx）",
        "非报关行：没有报关单号（内销/视同销售）": "内销/视同销售，本来就不产生报关单，不在按报关单匹配的范围",
        "非 MT 单：供应商自家编号（CY-/SM-）": "希蒙等供应商用自己的合同号（CY-25SM-S274），睿贝采购单表里没有对应单号",
    }
    for kind in list(FILLS) + [k for k in buckets if k not in FILLS]:
        if kind not in buckets:
            continue
        bucket = buckets[kind]
        summary.append(
            [
                kind,
                bucket["rows"],
                str(bucket["money"]),
                f"{bucket['money'] / total * 100:.1f}%",
                str(corroborated.get(kind, Decimal(0))) if kind in corroborated else "",
                notes.get(kind, ""),
            ]
        )
        fill = FILLS.get(kind)
        if fill:
            for cell in summary[summary.max_row]:
                cell.fill = fill
    summary.append(
        [
            "合计",
            len(records),
            str(total),
            "100.0%",
            str(sum(corroborated.values(), Decimal(0))),
            "",
        ]
    )
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for index, name in enumerate(summary[1], 1):
        summary.column_dimensions[summary.cell(row=1, column=index).column_letter].width = max(
            14, min(64, len(str(name)) * 2 + 10)
        )

    for title, kind in (
        ("非报关行明细", "非报关行：没有报关单号（内销/视同销售）"),
        ("供应商编号明细", "非 MT 单：供应商自家编号（CY-/SM-）"),
    ):
        sheet_ = workbook.create_sheet(title)
        sheet_.append(FIELDS)
        for cell in sheet_[1]:
            cell.font = Font(bold=True)
        for row in detail:
            if row["_kind"] != kind:
                continue
            sheet_.append([row.get(f) for f in FIELDS])
        for index, name in enumerate(FIELDS, 1):
            sheet_.column_dimensions[sheet_.cell(row=1, column=index).column_letter].width = max(
                12, min(36, len(str(name)) * 2 + 8)
            )
        sheet_.freeze_panes = "A2"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)

    print(f"飞书有采购金额 {len(records)} 行，合计 {total}")
    for kind, bucket in sorted(buckets.items(), key=lambda kv: -kv[1]["money"]):
        print(
            f"  {kind:<42} {bucket['rows']:>4} 行 {bucket['money']:>18}"
            f"（{bucket['money'] / total * 100:5.1f}%）"
            f" 已印证 {corroborated.get(kind, Decimal(0))}"
        )
    print("输出：", OUT)


if __name__ == "__main__":
    main()
