r"""逐条核对用的取证工具：一条命令把「这一个单元」的全部原始证据打出来。

用法：
    & '.\.venv\Scripts\python.exe' scripts\show_evidence.py 223120260003024940
    & '.\.venv\Scripts\python.exe' scripts\show_evidence.py 26MT-03T203Y-HX
    & '.\.venv\Scripts\python.exe' scripts\show_evidence.py 223120260003024940 310120260516694451

参数可以是**报关单号**或**采购单号**（可给多个），打印内容：

  ① 我方记录（`成本匹配_全部.xlsx`）：金额、分摊依据、判定、归因
  ② 拆单行（`拆单明细_全部.xlsx`）：出运金额(USD) / 出运采购金额(RMB) / 产品行数
  ③ 飞书行：采购金额 = 开票金额、报关金额、重量
  ④ 入库单：有效附件的实发金额、各批次区块、费用行，以及被丢弃的旧版
  ⑤ 睿贝：采购单金额、整单出运金额(RMB/USD)、按产品类型的出运金额

只用来看，不改任何数据。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook  # noqa: E402

from app.services.cost_match import shipment_money, unit_core  # noqa: E402
from app.services.erp_cache import CACHE_ROOT  # noqa: E402
from app.services.grn_extract import PARSED_ROOT, safe_name  # noqa: E402
from app.services.grn_select import erp_amount, erp_supplier, select  # noqa: E402

COST_DIR = PROJECT_ROOT / "outputs" / "cost_match"
SPLIT = PROJECT_ROOT / "outputs" / "shipments_split" / "拆单明细_全部.xlsx"
REF = PROJECT_ROOT / "data" / "cache" / "ref_2026_customs_full.json"
SHIPMENT_DETAILS = CACHE_ROOT / "details" / "shipments"


def sheet_rows(path: Path, sheet: str | None = None) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    target = workbook[sheet] if sheet else workbook.active
    data = list(target.iter_rows(values_only=True))
    header = [str(cell) for cell in data[0]]
    workbook.close()
    return [dict(zip(header, row)) for row in data[1:]]


def matched(row_value: str, wanted: set[str]) -> bool:
    return row_value.strip() in wanted or unit_core(row_value) in wanted


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        return
    wanted = {arg.strip() for arg in sys.argv[1:] if arg.strip()}
    wanted_cores = {unit_core(arg) for arg in wanted}

    cost_rows = sheet_rows(COST_DIR / "成本匹配_全部.xlsx")
    split_rows = sheet_rows(SPLIT)
    ref_rows = json.loads(REF.read_text(encoding="utf-8"))["records"]
    money = shipment_money()

    seen_codes: set[str] = set()
    for target in sorted(wanted):
        print("=" * 100)
        print(f"目标：{target}")
        hits = [
            row for row in cost_rows
            if str(row.get("报关单号") or "").strip() == target
            or str(row.get("采购单号") or "").strip() == target
        ]
        if not hits:
            print("  （成本匹配结果里没有这个报关单号/采购单号）")
        for row in hits:
            print(
                f"  ① 我方：{row['报关单号']}｜{row['合同号_1']}｜{row['供应商简称']}｜"
                f"{row['采购单号']}｜{row['产品类型']}｜判定={row['判定']}"
            )
            print(
                f"       金额={row['采购金额_我方']}｜飞书={row['采购金额_飞书']}｜"
                f"差异={row['差异']}｜依据={row['分摊依据']}"
            )
            if row.get("异常原因"):
                print(f"       原因={row['异常原因']}｜{row['归因说明']}")
            picked = str(row["采购单号"]).strip()
            if picked and picked not in seen_codes:
                seen_codes.add(picked)

        for row in split_rows:
            if str(row.get("报关单号") or "").strip() != target:
                continue
            print(
                f"  ② 拆单：{row['合同号_1']}｜{row['供应商简称']}｜{row['采购单号']}｜"
                f"{row['产品类型']}｜USD={row['出运金额合计']}｜RMB={row['出运采购金额合计']}｜"
                f"报关金额={row['报关金额']}｜客户费用={row['客户费用分摊']}｜行数={row['产品行数']}"
            )

        for record in ref_rows:
            decl = str(record.get("报关单号") or "").strip()
            code = str(record.get("合同号（应收表格）") or record.get("合同号_1") or "")
            if decl != target and not matched(code, wanted_cores):
                continue
            print(
                f"  ③ 飞书：{decl}｜{code}｜{record.get('供应商简称')}｜{record.get('报关品名')}｜"
                f"采购金额={record.get('采购金额')}｜开票金额={record.get('开票金额')}｜"
                f"报关金额={record.get('报关金额')}｜重量={record.get('报关重量')}"
            )

    for code in sorted(seen_codes):
        print("=" * 100)
        print(f"采购单 {code}")
        result = select(code)
        print(
            f"  ⑤ 睿贝：供应商={erp_supplier(code)}｜采购金额={erp_amount(code)}｜"
            f"整单出运RMB={money.rmb_by_po.get(code, '—')}｜整单出运USD={money.usd_by_po.get(code, '—')}"
        )
        for (purchase, kind), value in sorted(money.rmb_by_kind.items()):
            if purchase == code:
                print(f"       按产品类型 {kind} = {value}")
        print(f"  ④ 入库单：有效附件 {len(result['kept'])} 份，丢弃 {len(result['dropped'])} 份")
        for item in result["kept"]:
            path = PARSED_ROOT / safe_name(code) / (safe_name(item["file"], "file") + ".json")
            payload = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
            blocks = "、".join(
                f"{block.get('label')}={block.get('amount')}"
                for block in payload.get("settled_blocks") or []
                if isinstance(block, dict)
            )
            fees = payload.get("fee_by_column") or {}
            fee_text = "、".join(f"{key}={value}" for key, value in fees.items())
            print(
                f"       保留 {item['original']}｜来源={payload.get('_how')}｜实发={item['amount']}"
                f"｜区块[{blocks}]｜费用行[{fee_text}]"
            )
        for item in result["dropped"]:
            print(f"       丢弃 {item['original']}")
        for note in result["issues"]:
            print(f"       留痕：{note[:160]}")

    # 出运单明细里挂在这些采购单下的产品行（看数量/单价/金额）
    if seen_codes and SHIPMENT_DETAILS.exists():
        print("=" * 100)
        print("出运产品行（只列挂在本采购单下的行）")
        for path in sorted(SHIPMENT_DETAILS.glob("*.json")):
            payload = json.loads(path.read_text(encoding="utf-8"))
            rows = [
                row for row in payload.get("productList") or []
                if str(row.get("采购订单号") or "").strip() in seen_codes
            ]
            if not rows:
                continue
            print(f"  —— {path.name}")
            for row in rows:
                print(
                    f"     采购订单={row.get('采购订单号')}｜{row.get('海关商品（中文）')}｜"
                    f"数量={row.get('出运数量')}｜净重={row.get('每个净重')}｜"
                    f"USD={row.get('出运金额')}｜RMB={row.get('出运采购金额(RMB)')}"
                )
            for key in ("expenseList", "purchaseExpenseList"):
                if payload.get(key):
                    print(f"     {key}={json.dumps(payload[key], ensure_ascii=False)}")


if __name__ == "__main__":
    main()
