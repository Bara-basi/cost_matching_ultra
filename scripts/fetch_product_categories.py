r"""按产品编码抓睿贝商品资料（含 `类别名称`），供拆单回填 `产品类型`。

落盘：`.cache/erp/details/products/<safe(产品编码)>.json`（原始 product.find 返回）。
断点续跑：已缓存的产品编码直接跳过（`--force` 可强制重抓）。

范围（`--scope`）：
  needed（默认）  报关单解析结果涉及到的出运单里的产品编码 + 所有缺「海关商品（中文）」的行
  empty           只有缺「海关商品（中文）」的行
  all             出运明细里出现过的全部产品编码

用法：
  & '.\.venv\Scripts\python.exe' scripts\fetch_product_categories.py
  & '.\.venv\Scripts\python.exe' scripts\fetch_product_categories.py --scope all --pause 1.0
"""
from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from app.services.erp_cache import CACHE_ROOT  # noqa: E402
from app.services.product_master import product_path  # noqa: E402
from app.services.feishu_client import get_config  # noqa: E402
from scripts.mcp_erp import McpClient  # noqa: E402

DETAIL_DIR = CACHE_ROOT / "details" / "shipments"
PARSE_XLSX = PROJECT_ROOT / "outputs" / "customs_parse" / "报关单解析结果_出口退税联.xlsx"


def iter_shipment_rows():
    for path in sorted(DETAIL_DIR.glob("*.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except ValueError:
            continue
        invoice = str(payload.get("invoiceCode") or "").strip()
        for row in payload.get("productList") or []:
            yield invoice, row


def codes_for_scope(scope: str) -> list[str]:
    empty_codes: list[str] = []
    all_codes: list[str] = []
    seen_all: set[str] = set()
    seen_empty: set[str] = set()
    for _invoice, row in iter_shipment_rows():
        code = str(row.get("产品编码") or "").strip()
        if not code:
            continue
        if code not in seen_all:
            seen_all.add(code)
            all_codes.append(code)
        if not str(row.get("海关商品（中文）") or "").strip() and code not in seen_empty:
            seen_empty.add(code)
            empty_codes.append(code)
    if scope == "empty":
        return empty_codes
    if scope == "all":
        return all_codes

    # needed：解析结果涉及的出运单里的产品编码 + 全部缺名行的产品编码
    from openpyxl import load_workbook
    from app.services.shipment_detail import shipment_invoices_for_contract

    contracts: list[str] = []
    if PARSE_XLSX.exists():
        workbook = load_workbook(PARSE_XLSX, read_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        header = [str(cell) for cell in rows[0]]
        position = header.index("合同号_1")
        seen: set[str] = set()
        for row in rows[1:]:
            contract = str(row[position] or "").strip()
            if contract and contract not in seen:
                seen.add(contract)
                contracts.append(contract)
    invoices: set[str] = set()
    for contract in contracts:
        for invoice in shipment_invoices_for_contract(contract):
            invoices.add(str(invoice).strip().upper())
    # 先把「本轮解析结果真正会用到」的编码排在前面（优先出结果），再补其余缺名编码
    wanted: list[str] = []
    seen: set[str] = set()
    for invoice, row in iter_shipment_rows():
        if str(invoice or "").strip().upper() not in invoices:
            continue
        code = str(row.get("产品编码") or "").strip()
        if code and code not in seen:
            seen.add(code)
            wanted.append(code)
    for code in empty_codes:
        if code not in seen:
            seen.add(code)
            wanted.append(code)
    return wanted


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--token", default="")
    parser.add_argument("--scope", default="needed", choices=("needed", "empty", "all"))
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--pause", type=float, default=1.0)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    codes = codes_for_scope(args.scope)
    if args.limit:
        codes = codes[: args.limit]
    todo = [code for code in codes if args.force or not product_path(code).exists()]
    print(f"范围={args.scope} 目标产品编码={len(codes)} 待抓={len(todo)}", flush=True)
    if not todo:
        print("已全部缓存，无需抓取。", flush=True)
        return

    client = McpClient(args.token or get_config("ERP_API_KEY"))
    client.initialize()
    done = failed = 0
    for index, code in enumerate(todo, 1):
        try:
            text = client.call_text(
                "product.find", {"productQueryType": "productCode", "productQueryValue": code}
            )
            payload = json.loads(text)
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"  [{index}/{len(todo)}] {code} ERROR {exc}", flush=True)
            time.sleep(2.0)
            continue
        path = product_path(code)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
        done += 1
        if not payload.get("success"):
            print(f"  [{index}/{len(todo)}] {code} 未命中：{payload.get('msg')}", flush=True)
        time.sleep(args.pause)
        if index % 25 == 0 or index == len(todo):
            print(
                f"  [{index}/{len(todo)}] 已抓={done} 失败={failed} 暂停={args.pause}s",
                flush=True,
            )
    print(json.dumps({"done": done, "failed": failed}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
