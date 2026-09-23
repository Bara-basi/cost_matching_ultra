"""把拆单异常按「是否本项目范围」重新归类，输出待解决清单。"""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.services.scope import out_of_scope_reason  # noqa: E402

OUT_DIR = PROJECT_ROOT / "outputs" / "shipments_split"
PARSE_XLSX = PROJECT_ROOT / "outputs" / "customs_parse" / "报关单解析结果_出口退税联.xlsx"
REPORT = PROJECT_ROOT / ".cache" / "erp" / "reports"


def main() -> None:
    # 目标文件可能被执行占用，此时新版会带时间戳后缀：取最新的一份
    candidates = sorted(
        OUT_DIR.glob("拆单结果_异常*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True
    )
    source = candidates[0]
    print(f"读取异常表：{source.name}")
    workbook = load_workbook(source, read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    header = [str(x) for x in rows[0]]
    index = {name: pos for pos, name in enumerate(header)}
    records = [dict(zip(header, r)) for r in rows[1:]]

    parsed = load_workbook(PARSE_XLSX, read_only=True)
    p_rows = list(parsed.active.iter_rows(values_only=True))
    p_header = [str(x) for x in p_rows[0]]
    p_idx = {n: p for p, n in enumerate(p_header)}
    my_decls = {str(r[p_idx["报关单号"]] or "") for r in p_rows[1:]}

    buckets: dict[str, list[dict]] = {}

    def put(bucket: str, row: dict) -> None:
        buckets.setdefault(bucket, []).append(row)

    for row in records:
        result = str(row.get("比对结果") or "")
        contract = str(row.get("合同号_1") or "")
        decl = str(row.get("报关单号") or "")
        if out_of_scope_reason(contract):
            put("A. 超出范围：历史遗留合同号（CY/SM/SP）", row)
        elif result == "未拆出" and decl not in my_decls:
            put("B. 超出范围：飞书有此报关单，我没有对应 PDF", row)
        elif result == "未拆出":
            put("C. 待解决：有 PDF 但没拆出", row)
        elif result.startswith("无出运产品行"):
            put("C. 待解决：找不到出运产品行", row)
        elif "多解" in result or "不唯一" in result:
            put("C. 待解决：候选不唯一", row)
        elif result.startswith("不一致"):
            put("C. 待解决：供应商集合不一致", row)
        elif result.startswith("ERP数据异常"):
            put("C. 待解决：ERP数据异常（供应商与产品不符）", row)
        elif "无精确解" in result or "无解" in result or "未找到" in result or "跳过精确求解" in result:
            put("C. 待解决：子集和无解", row)
        else:
            put("C. 待解决：其它", row)

    out = [f"异常总数={len(records)}", ""]
    for name, items in sorted(buckets.items()):
        out.append(f"{name}: {len(items)}")
    out.append("")
    out.append("—— C 类（待解决）样例 ——")
    for row in (buckets.get("C. 待解决：有 PDF 但没拆出") or [])[:10]:
        out.append(f"    {row.get('报关单号')} {row.get('合同号_1')} {row.get('报关品名')} 飞书={row.get('飞书供应商简称')}")
    for row in (buckets.get("C. 待解决：供应商集合不一致") or [])[:10]:
        out.append(f"    {row.get('报关单号')} {row.get('合同号_1')} 系统={row.get('系统供应商简称')} 飞书={row.get('飞书供应商简称')}")

    # 输出待解决清单 Excel
    pending = [r for k, v in buckets.items() if k.startswith("C.") for r in v]
    wb = Workbook()
    ws = wb.active
    ws.title = "待解决"
    ws.append(["分类"] + header)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(bold=True)
    for name, items in sorted(buckets.items()):
        if not name.startswith("C."):
            continue
        for row in items:
            ws.append([name] + [row.get(c, "") for c in header])
    ws.freeze_panes = "A2"
    wb.save(OUT_DIR / "异常_待解决.xlsx")

    out.append("")
    out.append(f"待解决合计={len(pending)}（已输出 outputs/shipments_split/异常_待解决.xlsx）")
    REPORT.mkdir(parents=True, exist_ok=True)
    (REPORT / "exception_triage.txt").write_text("\n".join(out), encoding="utf-8")
    print("\n".join(out[:10]))


if __name__ == "__main__":
    main()
