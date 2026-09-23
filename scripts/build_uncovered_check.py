"""缺报关单的行,金额能不能采信？——用入库单余额反查财务填的采购金额。

一个（订单核心 × 供应商）下,入库单实发总额应当正好盖住它所有报关行的采购金额：

    未分配余额   = 入库单实发总额 − 我方已匹配各行合计
    飞书缺单行   = 飞书该组合全部行 − 我方已匹配的飞书行

判定：
  * **可采信**  ：两者相等（±1%）——财务给缺单行填的钱正好是入库单剩下的那块；
  * **飞书多**  ：飞书 > 余额——财务多算，或我方还缺入库单，逐条列出；
  * **飞书少**  ：飞书 < 余额——该采购单还没出完（货在库里、报关单未产生），正常；
  * **无入库单**：我方根本没有这张采购单的入库单（多为供应链代采、ERP 无凭证）。

输出：outputs/cost_match/缺单行金额核对.xlsx（明细 + 汇总两张表）
"""
from __future__ import annotations

import json
import re
import sys
from decimal import Decimal
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.services.cost_match import unit_core  # noqa: E402
from app.services.erp_cache import CACHE_ROOT, read_jsonl  # noqa: E402
from app.services.grn_extract import PARSED_ROOT  # noqa: E402
from app.services.grn_select import select  # noqa: E402
from app.services.supplier_names import SupplierNames  # noqa: E402

OUT = PROJECT_ROOT / "outputs" / "cost_match" / "缺单行金额核对.xlsx"
MATCH = PROJECT_ROOT / "outputs" / "cost_match" / "采购金额_匹配.xlsx"
REF = PROJECT_ROOT / "data" / "cache" / "ref_2026_customs_full.json"
NOTES = PROJECT_ROOT / "data" / "reference" / "uncovered_notes.json"
FILLS = {
    "可采信": PatternFill("solid", fgColor="E2EFDA"),
    "飞书多（需查）": PatternFill("solid", fgColor="FFC7CE"),
    "飞书少（未出完）": PatternFill("solid", fgColor="FCE4D6"),
    "无入库单": PatternFill("solid", fgColor="EDEDED"),
}


def dec(value) -> Decimal:
    if value in (None, ""):
        return Decimal(0)
    try:
        return Decimal(str(value).replace(",", "").strip())
    except Exception:
        return Decimal(0)


def norm(text: str) -> str:
    return str(text or "").replace(" ", "").replace("_", "").upper()


def audit_core(code: str) -> str:
    """比对用的订单核心：再剥掉尾部的工厂后缀（2-5 个字母）。

    `unit_core` 对标准单号（26MT-02N182-JX）已经剥了，但对 `26MT-DP002-XMLS`
    这类非标准单号不剥——飞书那边写的是 `26MT-DP002`，不补一刀就永远配不上。
    """
    return re.sub(r"-[A-Z]{2,5}$", "", unit_core(code))


def main() -> None:
    names = SupplierNames()
    purchases = {
        norm(str(row.get("purchase_code") or "")): str(row.get("supplierName") or "")
        for row in read_jsonl(CACHE_ROOT / "purchases" / "purchases.jsonl")
    }

    # 我方：每张采购单的有效入库单实发合计 → 按（核心 × 供应商）汇总
    po_totals: dict[tuple[str, str], Decimal] = {}
    for folder in sorted(p for p in PARSED_ROOT.iterdir() if p.is_dir()):
        code = folder.name
        result = select(code)
        total = sum((dec(item.get("amount")) for item in result["kept"]), Decimal(0))
        if not total:
            continue
        raw_supplier = purchases.get(norm(code), "")
        key = (audit_core(code), names.short(raw_supplier) if raw_supplier else "")
        po_totals[key] = po_totals.get(key, Decimal(0)) + total

    # 我方：已匹配到报关单的单元
    book = load_workbook(MATCH, read_only=True, data_only=True)
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(c) for c in rows[0]]
    units = [dict(zip(header, row)) for row in rows[1:]]
    book.close()
    mine: dict[tuple[str, str], dict] = {}
    for unit in units:
        key = (
            audit_core(unit.get("采购单核心") or unit.get("采购单号")),
            str(unit.get("供应商简称") or ""),
        )
        entry = mine.setdefault(key, {"decls": set(), "amount": Decimal(0)})
        entry["decls"].add(str(unit.get("报关单号") or ""))
        entry["amount"] += dec(unit.get("采购金额_我方"))

    records = json.loads(REF.read_text(encoding="utf-8")).get("records") or []
    notes = {}
    if NOTES.exists():
        notes = (json.loads(NOTES.read_text(encoding="utf-8")).get("notes") or {})
    ref_rows: dict[tuple[str, str], list[dict]] = {}
    for record in records:
        core = audit_core(record.get("合同号（应收表格）") or record.get("合同号_1"))
        if core:
            ref_rows.setdefault((core, str(record.get("供应商简称") or "").strip()), []).append(record)

    ledger: list[list] = []
    stats: dict[str, list] = {}
    for key, entries in sorted(ref_rows.items()):
        total = sum((dec(r.get("采购金额")) for r in entries), Decimal(0))
        entry = mine.get(key)
        matched_decls = entry["decls"] if entry else set()
        matched = sum(
            (dec(r.get("采购金额")) for r in entries if str(r.get("报关单号") or "") in matched_decls),
            Decimal(0),
        )
        unmatched = total - matched
        if unmatched <= Decimal("0.05"):
            continue
        po_total = po_totals.get(key, Decimal(0))
        our_amount = entry["amount"] if entry else Decimal(0)
        left = po_total - our_amount
        if po_total <= 0:
            kind = "无入库单"
            detail = "我方没有这张采购单的入库单（多为供应链代采 / ERP 无凭证）"
        else:
            tol = max(Decimal(1), po_total * Decimal("0.01"))
            if abs(left - unmatched) <= tol:
                kind = "可采信"
                detail = f"入库单余额 {left} = 飞书缺单行 {unmatched}"
            elif unmatched > left:
                kind = "飞书多（需查）"
                detail = f"飞书比入库单余额多 {unmatched - left}（飞书 {unmatched} / 余额 {left}）"
            else:
                kind = "飞书少（未出完）"
                detail = f"入库单余额比飞书多 {left - unmatched}（货在库里、报关单还没产生）"
        extra = notes.get(f"{key[0]}|{key[1]}")
        if extra and kind != "可采信":
            detail = f"【已人工归因】{extra}｜原判：{detail}"
        bucket = stats.setdefault(kind, [0, Decimal(0), Decimal(0)])
        bucket[0] += 1
        bucket[1] += unmatched
        bucket[2] += left
        ledger.append(
            [
                key[0],
                key[1],
                len(entries),
                str(total),
                str(matched),
                str(unmatched),
                str(po_total) if po_total else "",
                str(our_amount) if our_amount else "",
                str(left) if po_total else "",
                kind,
                detail,
            ]
        )

    order = {"可采信": 0, "飞书多（需查）": 1, "飞书少（未出完）": 2, "无入库单": 3}
    ledger.sort(key=lambda row: (order.get(row[9], 9), -dec(row[5])))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "缺单行核对"
    sheet.append(
        [
            "订单核心",
            "供应商",
            "飞书行数",
            "飞书合计",
            "其中我方已匹配",
            "缺单行金额",
            "入库单实发总额",
            "我方已分配",
            "未分配余额",
            "判定",
            "说明",
        ]
    )
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in ledger:
        sheet.append(row)
        fill = FILLS.get(str(row[9]))
        if fill:
            for cell in sheet[sheet.max_row]:
                cell.fill = fill
    for index, name in enumerate(sheet[1], 1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            14, min(64, len(str(name)) * 2 + 8)
        )
    sheet.freeze_panes = "A2"

    summary = workbook.create_sheet("汇总")
    summary.append(["判定", "组合数", "涉及飞书金额", "涉及入库单余额"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for kind in order:
        if kind not in stats:
            continue
        count, money, left = stats[kind]
        summary.append([kind, count, str(money), str(left)])
    for index, name in enumerate(summary[1], 1):
        summary.column_dimensions[summary.cell(row=1, column=index).column_letter].width = max(
            16, len(str(name)) * 2 + 8
        )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)

    print(f"有缺单行的组合 {len(ledger)} 个")
    for kind in order:
        if kind in stats:
            count, money, left = stats[kind]
            print(f"  {kind:<16} {count:>4} 个，飞书金额 {money}，入库单余额 {left}")
    print("输出：", OUT)


if __name__ == "__main__":
    main()
