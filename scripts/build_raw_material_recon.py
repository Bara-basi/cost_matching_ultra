"""原料购买家族三方对账单：睿贝采购单 × 我们的入库单 × 飞书行。

这些合同（08C020/08C652/08C577/08C596…）是麦金/瑞浦/久鸿的原料代采，
一个合同号下挂多家供应商、多批次，飞书行又没有报关单号，是剩余金额最集中的一块。

输出：outputs/cost_match/原料家族对账.xlsx（采购单与入库单 / 飞书行 / 按供应商小计）
"""
from __future__ import annotations

import collections
import json
import re
import sys
from decimal import Decimal
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.services.cost_match import unit_core  # noqa: E402
from app.services.erp_cache import CACHE_ROOT, read_jsonl  # noqa: E402
from app.services.grn_extract import PARSED_ROOT, safe_name  # noqa: E402
from app.services.grn_select import select  # noqa: E402
from app.services.supplier_names import SupplierNames  # noqa: E402

OUT = PROJECT_ROOT / "outputs" / "cost_match" / "原料家族对账.xlsx"
REF = PROJECT_ROOT / "data" / "cache" / "ref_2026_customs_full.json"
UNCOVERED_NOTES = PROJECT_ROOT / "data" / "reference" / "uncovered_notes.json"
FILL = PatternFill("solid", fgColor="FCE4D6")


def dec(value) -> Decimal:
    try:
        return Decimal(str(value or 0).replace(",", ""))
    except Exception:
        return Decimal(0)


def norm(text: str) -> str:
    return str(text or "").replace(" ", "").replace("_", "").upper()


def audit_core(code: str) -> str:
    """订单核心：剥掉尾部工厂后缀（DP 系列单号 unit_core 不剥，飞书却写不带后缀的）。"""
    return re.sub(r"-[A-Z]{2,5}$", "", unit_core(code))


def main() -> None:
    names = SupplierNames()
    records = json.loads(REF.read_text(encoding="utf-8")).get("records") or []

    # 需要看的合同：飞书里「非报关行」且未能印证 + 已归因的「飞书多」
    targets: set[str] = set()
    for record in records:
        contract = str(record.get("合同号（应收表格）") or record.get("合同号_1") or "")
        if record.get("采购金额") and not str(record.get("报关单号") or "").strip():
            targets.add(audit_core(contract))
    notes = {}
    if UNCOVERED_NOTES.exists():
        notes = json.loads(UNCOVERED_NOTES.read_text(encoding="utf-8")).get("notes") or {}
        for key in notes:
            core = str(key).split("|")[0]
            if core:
                targets.add(core)
    targets.discard("")

    purchases = list(read_jsonl(CACHE_ROOT / "purchases" / "purchases.jsonl"))

    po_rows: list[list] = []
    po_totals: dict[tuple[str, str], Decimal] = collections.defaultdict(Decimal)
    for core in sorted(targets):
        for row in purchases:
            code = str(row.get("purchase_code") or "")
            if audit_core(code) != core:
                continue
            supplier = names.short(str(row.get("supplierName") or ""))
            result = select(code)
            kept = sum((dec(item.get("amount")) for item in result["kept"]), Decimal(0))
            po_totals[(core, supplier)] += kept
            po_rows.append(
                [
                    core,
                    code,
                    supplier,
                    str(row.get("amount") or ""),
                    str(row.get("grnQuantity") or ""),
                    str(row.get("lastGrnDate") or ""),
                    str(kept),
                    f"{len(result['kept'])}/{len(result['kept']) + len(result['dropped'])}",
                ]
            )
            for item in result["kept"]:
                path = PARSED_ROOT / safe_name(code) / (safe_name(item["file"], "file") + ".json")
                if not path.exists():
                    continue
                payload = json.loads(path.read_text(encoding="utf-8"))
                for block in payload.get("settled_blocks") or []:
                    if not isinstance(block, dict) or not block.get("amount"):
                        continue
                    po_rows.append(
                        [
                            core,
                            code,
                            f"    └ 区块 {block.get('label')}",
                            "",
                            "",
                            "",
                            str(dec(block.get("amount"))),
                            "",
                        ]
                    )

    ref_rows: list[list] = []
    ref_totals: dict[tuple[str, str], Decimal] = collections.defaultdict(Decimal)
    for record in sorted(
        records, key=lambda r: (unit_core(r.get("合同号（应收表格）") or ""), str(r.get("合同号（应收表格）")))
    ):
        core = audit_core(record.get("合同号（应收表格）") or record.get("合同号_1"))
        if core not in targets or not record.get("采购金额"):
            continue
        supplier = str(record.get("供应商简称") or "").strip()
        ref_totals[(core, supplier)] += dec(record.get("采购金额"))
        ref_rows.append(
            [
                core,
                str(record.get("合同号（应收表格）") or ""),
                supplier,
                str(record.get("采购金额") or ""),
                str(record.get("报关金额") or ""),
                str(record.get("报关重量") or ""),
                "有" if str(record.get("报关单号") or "").strip() else "无",
                str(record.get("销售方式") or ""),
                notes.get(f"{core}|{supplier}", ""),
            ]
        )

    summary: list[list] = []
    for key in sorted(set(po_totals) | set(ref_totals)):
        erp = sum(
            (
                dec(row.get("amount"))
                for row in purchases
                if audit_core(str(row.get("purchase_code") or "")) == key[0]
                and names.short(str(row.get("supplierName") or "")) == key[1]
            ),
            Decimal(0),
        )
        mine = po_totals.get(key, Decimal(0))
        ref = ref_totals.get(key, Decimal(0))
        diff = ref - mine
        if mine == 0 and ref == 0:
            continue
        if abs(diff) <= max(Decimal(1), mine * Decimal("0.01")):
            verdict = "一致"
        elif mine and ref < mine:
            verdict = f"飞书少 {mine - ref}（我方入库单里这部分还没出）"
        else:
            verdict = f"飞书多 {diff}（需业务对单）"
        summary.append(
            [key[0], key[1], str(erp), str(mine), str(ref), str(diff), verdict]
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "采购单与入库单"
    sheet.append(["合同核心", "睿贝采购单", "供应商", "睿贝金额", "入库数量", "最后入库日", "我方入库单合计", "保留/附件数"])
    for row in po_rows:
        sheet.append(row)
    sheet2 = workbook.create_sheet("飞书行")
    sheet2.append(["合同核心", "合同号", "供应商", "飞书采购金额", "报关金额", "报关重量", "报关单号", "销售方式", "已归因说明"])
    for row in ref_rows:
        sheet2.append(row)
    sheet3 = workbook.create_sheet("按供应商小计")
    sheet3.append(["合同核心", "供应商", "睿贝金额", "我方入库单", "飞书", "飞书−我方", "判定"])
    for row in summary:
        sheet3.append(row)
        if "需业务对单" in str(row[-1]):
            for cell in sheet3[sheet3.max_row]:
                cell.fill = FILL
    for target, widths in ((sheet, None), (sheet2, None), (sheet3, None)):
        for cell in target[1]:
            cell.font = Font(bold=True)
        for index in range(1, target.max_column + 1):
            target.column_dimensions[target.cell(row=1, column=index).column_letter].width = 18
        target.freeze_panes = "A2"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)

    print(f"涉及合同核心 {len(targets)} 个；采购单行 {len(po_rows)}；飞书行 {len(ref_rows)}；小计 {len(summary)}")
    for row in summary:
        if "需业务对单" in str(row[-1]):
            print(f"  ★ {row[0]:<12}{row[1]:<6} 睿贝 {row[2]:>14} 我方 {row[3]:>14} 飞书 {row[4]:>14} 差 {row[5]:>14}")
    print("输出：", OUT)


if __name__ == "__main__":
    main()
