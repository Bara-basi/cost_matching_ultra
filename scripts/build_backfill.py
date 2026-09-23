"""把匹配结果整理成「照着改飞书」的回填建议表。

一行 = 飞书里的一条（报关单 + 品名 + 供应商 + 采购单号）记录，
带上我方金额、差异、依据与建议动作。飞书里采购金额为空的行不处理。

输出：`outputs/cost_match/回填建议.xlsx`
"""
from __future__ import annotations

import json
import sys
from decimal import Decimal
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.services.contract_shipments import cores_of  # noqa: E402
from app.services.cost_match import dec, unit_core  # noqa: E402

OUT_DIR = PROJECT_ROOT / "outputs" / "cost_match"
REF_CACHE = PROJECT_ROOT / "data" / "cache" / "ref_2026_customs.json"
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFE699")


def po_core(code: str) -> str:
    return unit_core(code)


def read_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(cell) for cell in rows[0]]
    out = [dict(zip(header, row)) for row in rows[1:]]
    workbook.close()
    return out


def main() -> None:
    units = read_rows(OUT_DIR / "采购金额_单元差异.xlsx")
    unit_map = {
        (u["报关单号"], u["供应商简称"], u["采购单核心"]): u for u in units
    }
    details = read_rows(OUT_DIR / "采购金额_匹配.xlsx")
    review = read_rows(OUT_DIR / "待确认清单.xlsx") if (OUT_DIR / "待确认清单.xlsx").exists() else []
    review_map = {
        (r["报关单号"], r["供应商简称"], r["采购单核心"]): str(r.get("建议") or "")
        for r in review
    }
    basis_map: dict[tuple, list[str]] = {}
    for row in details:
        key = (row["报关单号"], row["供应商简称"], row["采购单核心"])
        basis_map.setdefault(key, []).append(str(row.get("分摊依据") or ""))

    ref = json.loads(REF_CACHE.read_text(encoding="utf-8"))["records"]
    grouped: dict[tuple, list[dict]] = {}
    for row in ref:
        amount = row.get("采购金额")
        if amount in (None, ""):
            continue
        key = (
            str(row.get("报关单号") or "").strip(),
            str(row.get("供应商简称") or "").strip(),
            po_core(str(row.get("采购单号") or "")),
        )
        grouped.setdefault(key, []).append(row)

    out: list[list] = []
    for key, rows in sorted(grouped.items()):
        unit = unit_map.get(key)
        if unit is None:
            continue
        mine = dec(unit["采购金额_我方"])
        theirs = sum((dec(r.get("采购金额")) for r in rows), Decimal(0))
        basis = "；".join(sorted(set(basis_map.get(key, []))))
        if theirs and abs(mine - theirs) <= max(Decimal("0.01"), theirs * Decimal("0.01")):
            action = "无需调整"
        elif review_map.get(key, "").startswith("已核实"):
            action = review_map[key]
        elif review_map.get(key, "").startswith("建议以我方为准"):
            action = f"建议改为 {mine}（{review_map[key]}）"
        elif len(rows) == 1 and mine > 0:
            action = f"待确认口径后决定（我方 {mine}，飞书 {theirs}）"
        elif mine > 0:
            action = f"本单元合计应为 {mine}（飞书现合计 {theirs}），请按品名/重量分摊"
        else:
            action = "需人工（我方无金额）"
        for row in rows:
            out.append(
                [
                    row.get("报关单号"),
                    row.get("报关品名"),
                    row.get("供应商简称"),
                    row.get("采购单号"),
                    str(row.get("采购金额")),
                    str(mine),
                    "" if not theirs else str(mine - theirs),
                    action,
                    unit.get("结论"),
                    basis,
                ]
            )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "回填建议"
    header = [
        "报关单号", "报关品名(飞书)", "供应商简称", "飞书采购单号", "飞书采购金额",
        "我方采购金额", "差异", "建议动作", "系统判定", "我方分摊依据",
    ]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in out:
        sheet.append(row)
        action = str(row[7])
        if action == "无需调整":
            fill = OK_FILL
        elif action.startswith(("需人工", "待确认")):
            fill = WARN_FILL
        else:
            fill = BAD_FILL
        for cell in sheet[sheet.max_row]:
            cell.fill = fill
    for index, name in enumerate(header, 1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            12, min(52, len(name) * 2 + 4)
        )
    sheet.freeze_panes = "A2"
    path = OUT_DIR / "回填建议.xlsx"
    workbook.save(path)
    change = sum(1 for row in out if str(row[7]).startswith("建议改为"))
    print(f"written {path} ({len(out)} rows)，其中「建议改为」{change} 行")

    # 最小改动清单：只留财务改飞书要用的列
    minimal = [
        [row[0], row[1], row[2], row[3], row[5], row[4], row[7]]
        for row in out
        if str(row[7]).startswith(("建议改为", "已核实"))
    ]
    book = Workbook()
    sheet = book.active
    sheet.title = "待改"
    head = ["报关单号", "报关品名", "供应商简称", "飞书采购单号", "建议采购金额", "飞书现值", "依据"]
    sheet.append(head)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in minimal:
        sheet.append(row)
    for index, name in enumerate(head, 1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            12, min(60, len(name) * 2 + 6)
        )
    sheet.freeze_panes = "A2"
    minimal_path = OUT_DIR / "待改清单_最小.xlsx"
    book.save(minimal_path)
    print(f"written {minimal_path} ({len(minimal)} rows)")


if __name__ == "__main__":
    main()
