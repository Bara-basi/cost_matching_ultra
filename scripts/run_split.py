"""对 PDF 解析结果执行拆单，输出正常/异常两份 Excel，并与飞书比对。

输出（outputs/shipments_split/）：
- 拆单结果_正常.xlsx
- 拆单结果_异常.xlsx
"""
from __future__ import annotations

import argparse
import json
import sys
from decimal import Decimal
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

from app.services.split_engine import DeclaredLine, SplitEngine  # noqa: E402
from app.services.scope import out_of_scope_reason  # noqa: E402

PARSE_XLSX = PROJECT_ROOT / "outputs" / "customs_parse" / "报关单解析结果_出口退税联.xlsx"
OUT_DIR = PROJECT_ROOT / "outputs" / "shipments_split"
CACHE = PROJECT_ROOT / "data" / "cache"

# 已知 PDF 解析不全导致单号错误的两条，直接过滤
BAD_CONTRACTS = {
    "26mt-03p200y-a&229y-a&245y-a&",
    "26mt-03p200y-a&229y-a&245y-a&262",
}

OUT_COLUMNS = (
    "来源文件",
    "合同号_1",
    "报关单号",
    "商品序号",
    "报关品名",
    "海关编码",
    "产品类型",
    "供应商",
    "采购单号",
    "出运单ID",
    "报关重量",
    "报关金额",
    "币种",
    "比对结果",
)

GROUP_COLUMNS = (
    "报关单号",
    "合同号_1",
    "报关品名",
    "海关编码",
    "系统供应商数",
    "系统供应商",
    "飞书供应商",
    "系统重量合计",
    "飞书重量合计",
    "比对结果",
)

ERROR_COLUMNS = (
    "来源文件",
    "合同号_1",
    "报关单号",
    "报关品名",
    "海关编码",
    "报关重量",
    "异常原因",
)

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
HEADER_FONT = Font(bold=True)


def flatten(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return " | ".join(filter(None, (flatten(v) for v in value)))
    if isinstance(value, dict):
        for key in ("text", "name", "value"):
            if key in value:
                return flatten(value[key])
    return ""


def read_declared() -> list[DeclaredLine]:
    from openpyxl import load_workbook

    workbook = load_workbook(PARSE_XLSX, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(x) for x in rows[0]]
    idx = {name: pos for pos, name in enumerate(header)}
    lines: list[DeclaredLine] = []
    seen: set[tuple] = set()
    for row in rows[1:]:
        def cell(name: str):
            pos = idx.get(name)
            return row[pos] if pos is not None and row[pos] is not None else ""

        contract = str(cell("合同号_1")).strip()
        if not contract or contract.lower() in BAD_CONTRACTS:
            continue
        # 二次拦截：历史遗留合同号（CY / xxSM / SP）不进成本匹配
        if out_of_scope_reason(contract):
            continue
        weight = str(cell("报关重量") or 0).replace(",", "") or "0"
        try:
            weight_value = Decimal(weight)
        except Exception:  # noqa: BLE001
            weight_value = Decimal(0)
        declared = str(cell("总价") or 0).replace(",", "") or "0"
        try:
            amount_value = Decimal(declared)
        except Exception:  # noqa: BLE001
            amount_value = Decimal(0)
        # 同一张报关单在 data/raw 下有重复 PDF 时会产生重复行：
        # 用「报关单号+合同号+品名+重量+金额」判重，只保留第一条
        source_file = str(cell("来源文件"))
        marker = (
            str(cell("报关单号")),
            contract,
            str(cell("报关品名")),
            weight,
            declared,
        )
        if marker in seen:
            continue
        seen.add(marker)
        lines.append(
            DeclaredLine(
                source_file=source_file,
                contract=contract,
                declaration_no=str(cell("报关单号")),
                product_name=str(cell("报关品名")),
                hs_code=str(cell("海关编码")),
                weight=weight_value,
                amount=amount_value,
                currency=str(cell("币种")),
                product_type=str(cell("产品类型") or ""),
            )
        )
    return lines


def write_sheet(rows: list[dict], path: Path, columns: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "拆单结果"
    columns = list(columns or OUT_COLUMNS)
    sheet.append(columns)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append([row.get(name, "") for name in columns])
    for column, name in zip(sheet.iter_cols(min_row=1, max_row=1), columns):
        sheet.column_dimensions[column[0].column_letter].width = max(10, min(30, len(name) * 2))
    sheet.freeze_panes = "A2"
    workbook.save(path)


def load_reference() -> dict[tuple[str, str], dict]:
    """飞书参考：报关单号 + 报关品名 -> {suppliers, weight, contract}。"""
    from collections import defaultdict

    path = CACHE / "records_ai.json"
    if not path.exists():
        return {}
    out: dict[tuple[str, str], dict] = defaultdict(
        lambda: {"suppliers": set(), "weight": Decimal(0), "contract": ""}
    )
    for record in json.loads(path.read_text(encoding="utf-8")):
        fields = record.get("fields", {})
        decl = flatten(fields.get("报关单号"))
        name = flatten(fields.get("报关品名"))
        contract = flatten(fields.get("合同号_1"))
        supplier = flatten(fields.get("供应商"))
        if not (decl and name):
            continue
        entry = out[(decl, name)]
        if supplier:
            entry["suppliers"].add(supplier)
        if not entry["contract"] and contract:
            entry["contract"] = contract
        weight = flatten(fields.get("报关重量"))
        if weight:
            try:
                entry["weight"] += Decimal(weight.replace(",", ""))
            except Exception:  # noqa: BLE001
                pass
    return out


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=0)
    args = parser.parse_args()

    lines = read_declared()
    if args.limit:
        lines = lines[: args.limit]
    print(f"待拆报关商品行={len(lines)}", flush=True)

    engine = SplitEngine()
    outcome = engine.split_many(lines)

    reference = load_reference()
    # 按「报关单号 + 报关品名」聚合后比对供应商集合（避免逐行错位比对）
    system_groups: dict[tuple[str, str], dict] = {}
    for row in outcome.rows:
        key = (row.declaration_no, row.product_name)
        group = system_groups.setdefault(
            key,
            {
                "报关单号": row.declaration_no,
                "合同号_1": row.contract,
                "报关品名": row.product_name,
                "海关编码": row.hs_code,
                "suppliers": set(),
                "weight": Decimal(0),
            },
        )
        if row.supplier:
            group["suppliers"].add(row.supplier)
        group["weight"] += row.weight

    normal: list[dict] = []
    abnormal: list[dict] = []
    matched = unmatched = 0
    for key, group in system_groups.items():
        ref = reference.get(key)
        expected = ref["suppliers"] if ref else None
        record = {
            "报关单号": group["报关单号"],
            "合同号_1": group["合同号_1"],
            "报关品名": group["报关品名"],
            "海关编码": group["海关编码"],
            "系统供应商数": len(group["suppliers"]),
            "系统供应商": "、".join(sorted(group["suppliers"])),
            "飞书供应商": "、".join(sorted(expected)) if expected else "(未填)",
            "系统重量合计": str(group["weight"]),
            "飞书重量合计": str(ref["weight"]) if ref else "",
        }
        if expected is None:
            record["比对结果"] = "飞书无此报关单"
            abnormal.append(record)
            unmatched += 1
            continue
        if not expected:
            record["比对结果"] = "飞书未填供应商"
            abnormal.append(record)
            unmatched += 1
            continue
        if group["suppliers"] == expected:
            record["比对结果"] = "供应商集合一致"
            normal.append(record)
            matched += 1
        else:
            record["比对结果"] = (
                f"供应商集合不一致: 系统缺={sorted(expected - group['suppliers'])} "
                f"系统多={sorted(group['suppliers'] - expected)}"
            )
            abnormal.append(record)
            unmatched += 1

    # 明细行输出（全部拆出的记录）
    detail_rows: list[dict] = []
    for row in outcome.rows:
        detail_rows.append(
            {
            "来源文件": row.source_file,
            "合同号_1": row.contract,
            "报关单号": row.declaration_no,
            "商品序号": "",
            "报关品名": row.product_name,
            "海关编码": row.hs_code,
            "产品类型": row.product_type,
            "供应商": row.supplier,
            "采购单号": row.purchase_code,
            "出运单ID": row.shipment_id,
            "报关重量": str(row.weight),
            "报关金额": str(row.purchase_amount),
            "币种": row.currency,
            "比对结果": "",
            }
        )
    for error in outcome.errors:
        abnormal.append({**error, "比对结果": "未拆分"})

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_sheet(detail_rows, OUT_DIR / "拆单明细_全部.xlsx", list(OUT_COLUMNS))
    write_sheet(normal, OUT_DIR / "拆单结果_正常.xlsx", list(GROUP_COLUMNS))
    write_sheet(
        abnormal,
        OUT_DIR / "拆单结果_异常.xlsx",
        list(GROUP_COLUMNS) + ["异常原因", "来源文件"],
    )
    summary = {
        "输入报关商品行": len(lines),
        "拆出记录": len(outcome.rows),
        "比对单元(报关单+品名)": len(system_groups),
        "正常": len(normal),
        "异常": len(abnormal),
        "供应商集合一致": matched,
        "供应商集合不一致或无对应": unmatched,
        "未拆分": len(outcome.errors),
    }
    (OUT_DIR / "拆单统计.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
