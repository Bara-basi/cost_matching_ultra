"""导出「附件名与采购单对不上」的清单（文件名、单号、上传顺序、处理结果）。

三类：
  1. 单号不符 —— 附件名里的采购单号与所在目录不一致；
  2. 供应商不符 —— 附件名点到别家工厂；
  3. 年份笔误 —— 只差 25MT/26MT 前缀（正常，仅登记）。

输出：outputs/grn_attachments/命名异常清单.xlsx
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from app.services.cost_match import unit_core  # noqa: E402
from app.services.grn_extract import PARSED_ROOT  # noqa: E402
from app.services.grn_select import (  # noqa: E402
    ORDER_KEY_RE,
    erp_amount,
    erp_supplier,
    load_items,
    order_key,
    same_order,
    select,
)

OUT = PROJECT_ROOT / "outputs" / "grn_attachments" / "命名异常清单.xlsx"


def head(key: str) -> str:
    return key.split("#")[0]


def main() -> None:
    rows_order: list[list] = []
    rows_factory: list[list] = []
    rows_year: list[list] = []
    for folder in sorted(p for p in PARSED_ROOT.iterdir() if p.is_dir()):
        code = folder.name
        items = load_items(code)
        if not items:
            continue
        result = select(code)
        kept = {item["original"] for item in result["kept"]}
        dropped = {item["original"] for item in result["dropped"]}
        target = order_key(code)
        target_base = head(target)
        for item in items:
            name = item.original
            matches = [m.group(0) for m in ORDER_KEY_RE.finditer(name)]
            if not matches:
                continue
            bases = {head(order_key(m)) for m in matches if order_key(m)}
            years = {re.match(r"(\d{2})", m).group(1) for m in matches if re.match(r"(\d{2})", m)}
            dir_year = re.match(r"(\d{2})", code)
            dir_year = dir_year.group(1) if dir_year else ""
            status = "已保留" if name in kept else ("已剔除" if name in dropped else "未处理")
            if item.other_order:
                kinds = []
                if target_base in bases:
                    kinds.append("同单号、不同工厂")
                else:
                    kinds.append("指向别的单号")
                rows_order.append(
                    [
                        code,
                        erp_supplier(code),
                        name,
                        item.order,
                        "、".join(sorted(bases)),
                        target_base,
                        "；".join(kinds),
                        status,
                    ]
                )
            elif item.foreign_factory:
                rows_factory.append(
                    [code, erp_supplier(code), name, item.order, "、".join(sorted(bases)), status]
                )
            elif target_base and target_base in bases and years and dir_year not in years:
                rows_year.append(
                    [code, erp_supplier(code), name, item.order, "、".join(sorted(years)), dir_year, status]
                )

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = [
        (
            "单号不符",
            ["采购单目录", "睿贝供应商", "附件原名", "上传序号", "附件里的单号", "目录单号", "判定", "处理结果"],
            rows_order,
        ),
        (
            "供应商不符",
            ["采购单目录", "睿贝供应商", "附件原名", "上传序号", "附件里的单号", "处理结果"],
            rows_factory,
        ),
        (
            "年份笔误",
            ["采购单目录", "睿贝供应商", "附件原名", "上传序号", "附件里的年份", "目录年份", "处理结果"],
            rows_year,
        ),
    ]
    for title, header, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for index, name in enumerate(header, 1):
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
                12, min(80, len(str(name)) * 2 + 8)
            )
        sheet.freeze_panes = "A2"
        print(f"{title}: {len(rows)} 行")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)
    print("输出：", OUT)


if __name__ == "__main__":
    main()
