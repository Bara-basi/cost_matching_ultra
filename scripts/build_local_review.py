r"""本地核对：只按**我方系统能确定的事实**判定每条拆单记录是否正常/异常。

本轮（2026-09-28）不再拿飞书当答案，所以这里不做任何飞书比对；判定只用：
  1. 报关单解析结果（PDF 原件事实：报关金额/报关重量/品名/海关编码）；
  2. 出运单明细（出运金额 / 出运采购金额(RMB) / 产品类型）；
  3. 采购单入库单实发（成本真值来源）。

输出 `outputs/local_review/`：
  - `本地核对_全部.xlsx`  每条拆单记录一行：判定 / 异常说明
  - `本地核对_异常.xlsx`  只保留异常
  - `本地核对_正常.xlsx`  只保留正常

用法：
  & '.\.venv\Scripts\python.exe' scripts\build_local_review.py
"""
from __future__ import annotations

import collections
import json
import sys
from decimal import Decimal
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.services.cost_match import dec, shipment_money  # noqa: E402
from app.services.erp_cache import CACHE_ROOT  # noqa: E402
from scripts.build_cost_match import (  # noqa: E402
    allocate_all,
    quantity_check,
)

SPLIT = PROJECT_ROOT / "outputs" / "shipments_split" / "拆单明细_全部.xlsx"
PARSE_XLSX = PROJECT_ROOT / "outputs" / "customs_parse" / "报关单解析结果_出口退税联.xlsx"
OUT_DIR = PROJECT_ROOT / "outputs" / "local_review"
EQUAL = Decimal("1")


def read_sheet(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    header = [str(cell) for cell in rows[0]]
    return [dict(zip(header, row)) for row in rows[1:]]


def parsed_weights() -> dict[str, Decimal]:
    """报关单号 → 报关重量（同一张单重复行按「品名+重量+金额」去重，与拆单同口径）。"""
    out: dict[str, Decimal] = collections.defaultdict(Decimal)
    seen: set[tuple] = set()
    for row in read_sheet(PARSE_XLSX):
        decl = str(row.get("报关单号") or "").strip()
        if not decl:
            continue
        marker = (
            decl,
            str(row.get("合同号_1") or ""),
            str(row.get("报关品名") or ""),
            str(row.get("报关重量") or ""),
            str(row.get("总价") or ""),
        )
        if marker in seen:
            continue
        seen.add(marker)
        out[decl] += dec(row.get("报关重量"))
    return dict(out)


def main() -> None:
    rows = read_sheet(SPLIT)
    print(f"拆单记录 {len(rows)} 条", flush=True)

    costs = allocate_all(rows)
    parsed_weight = parsed_weights()
    money = shipment_money()

    decl_weight: dict[str, Decimal] = collections.defaultdict(Decimal)
    decl_declared: dict[str, Decimal] = collections.defaultdict(Decimal)
    derived: list[dict] = []
    for row, cost in zip(rows, costs):
        decl = str(row.get("报关单号") or "")
        weight = dec(row.get("报关重量"))
        decl_weight[decl] += weight
        decl_declared[decl] += dec(row.get("报关金额"))
        derived.append(
            {
                **row,
                "采购金额_我方": cost.get("amount", ""),
                "分摊依据": cost.get("basis", ""),
                "采购单实发总额": cost.get("total", ""),
            }
        )

    # 出运数量 > 入库数量（本地可判定的组级异常）——明细口径直接取自 allocate_all 的 detail
    po_rows: dict[str, dict] = {}
    for cost in costs:
        code = str(cost.get("code") or "").strip()
        if not code or code in po_rows:
            continue
        detail = cost.get("detail") or {}
        po_rows[code] = {
            "采购单号": code,
            "数量(按单位族)": dict(detail.get("qty_by_family") or {}),
            "入库块数量": str(detail.get("block_qty") or ""),
            "入库单": "、".join(detail.get("kept_names") or []),
        }
    qty_anomalies = quantity_check(po_rows)

    out_rows: list[dict] = []
    for row in derived:
        issues: list[str] = []
        product_type = str(row.get("产品类型") or "").strip()
        if not product_type:
            issues.append("产品类型缺失（商品资料无类别且报关品名兜底也空）")
        rmb = dec(row.get("出运采购金额合计"))
        if rmb <= 0:
            issues.append("睿贝未定价（出运采购金额(RMB)=0）")
        total = dec(row.get("采购单实发总额"))
        if total <= 0:
            issues.append("未取得入库单实发金额（该采购单没有可用入库单）")
        basis = str(row.get("分摊依据") or "")
        if "空白" in basis or "未定价" in basis or "没有" in basis:
            issues.append(f"成本未落实：{basis}")
        weight = dec(row.get("报关重量"))
        if weight <= 0:
            issues.append("报关重量缺失/为 0")
        declared = dec(row.get("报关金额"))
        usd = dec(row.get("出运金额合计"))
        fee = dec(row.get("客户费用分摊"))
        if abs(usd + fee - declared) > Decimal("0.01"):
            issues.append(f"报关金额恒等式不成立（{usd}+{fee}≠{declared}）")
        po = str(row.get("采购单号") or "").strip()
        entry = qty_anomalies.get(po)
        if entry:
            issues.append(
                f"出运数量 > 入库数量（出运 {entry['出运数量']} / 入库 {entry['入库数量']}，"
                f"出运单 {entry['出运单']}）"
            )
        po_total = money.rmb_by_po.get(po, Decimal(0))
        if po_total > 0:
            ours = sum(
                (dec(r.get("出运采购金额合计")) for r in rows
                 if str(r.get("采购单号") or "").strip() == po),
                Decimal(0),
            )
            if ours + EQUAL < po_total:
                issues.append(
                    f"该采购单报关单不齐（我方货值 {ours} < 睿贝整单 {po_total}）"
                )
        out_rows.append(
            {
                "判定": "异常" if issues else "正常",
                "异常说明": "；".join(issues),
                "报关单号": row.get("报关单号"),
                "合同号_1": row.get("合同号_1"),
                "报关品名": row.get("报关品名"),
                "产品类型": product_type,
                "供应商简称": row.get("供应商简称"),
                "采购单号": po,
                "出运金额合计": row.get("出运金额合计"),
                "出运采购金额合计": row.get("出运采购金额合计"),
                "报关金额": row.get("报关金额"),
                "客户费用分摊": row.get("客户费用分摊"),
                "报关重量": row.get("报关重量"),
                "重量分摊依据": row.get("重量分摊依据"),
                "采购金额_我方": row.get("采购金额_我方"),
                "分摊依据": row.get("分摊依据"),
            }
        )

    # 报关单级：各行报关重量之和是否等于报关单原件重量
    for row in out_rows:
        decl = str(row.get("报关单号") or "")
        want = parsed_weight.get(decl)
        if want is None:
            continue
        got = decl_weight.get(decl, Decimal(0))
        if abs(got - want) > Decimal("0.01"):
            row["判定"] = "异常"
            note = f"报关重量合计≠报关单原件（我方 {got} / 原件 {want}）"
            row["异常说明"] = f"{row['异常说明']}；{note}" if row["异常说明"] else note

    normal = [row for row in out_rows if row["判定"] == "正常"]
    bad = [row for row in out_rows if row["判定"] == "异常"]
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    columns = [
        "判定", "异常说明", "报关单号", "合同号_1", "报关品名", "产品类型", "供应商简称",
        "采购单号", "出运金额合计", "出运采购金额合计", "报关金额", "客户费用分摊",
        "报关重量", "重量分摊依据", "采购金额_我方", "分摊依据",
    ]
    for name, data in (
        ("本地核对_全部.xlsx", out_rows),
        ("本地核对_正常.xlsx", normal),
        ("本地核对_异常.xlsx", bad),
    ):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "本地核对"
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
        for row in data:
            sheet.append([row.get(column, "") for column in columns])
        sheet.freeze_panes = "A2"
        workbook.save(OUT_DIR / name)
        print(f"  written {name} ({len(data)})")

    # 汇总
    reasons: collections.Counter = collections.Counter()
    for row in bad:
        for part in str(row["异常说明"]).split("；"):
            if part:
                reasons[part.split("（")[0]] += 1
    by_po: collections.Counter = collections.Counter()
    for row in bad:
        by_po[str(row.get("采购单号") or "")] += 1
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总"
    sheet.append(["项目", "数值"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.append(["拆单记录", len(out_rows)])
    sheet.append(["正常", len(normal)])
    sheet.append(["异常", len(bad)])
    sheet.append([])
    sheet.append(["异常原因", "条数"])
    for name, count in reasons.most_common():
        sheet.append([name, count])
    sheet.append([])
    sheet.append(["异常采购单", "条数"])
    for name, count in by_po.most_common():
        sheet.append([name, count])
    workbook.save(OUT_DIR / "本地核对_汇总.xlsx")
    print("  written 本地核对_汇总.xlsx")
    print(json.dumps(
        {"记录": len(out_rows), "正常": len(normal), "异常": len(bad),
         "异常原因": dict(reasons.most_common())},
        ensure_ascii=False,
    ))


if __name__ == "__main__":
    main()
